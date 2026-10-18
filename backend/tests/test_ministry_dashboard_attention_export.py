"""C14 and C15 — an actionable Ministry dashboard, and its masked background stage-data export.

C14 (owner decision, 2026-10-17): clickable cards (awaiting approval, returned, overdue, setup
incomplete), a state breakdown and the sanctioned rupees, each counted over the caller's scope and
each narrowing the SAME list. C15: one spreadsheet with a sheet per record type, identity numbers
masked, built by the queue worker.

The pure half needs no database; the rest follows ``tests/test_sanction_orders.py``'s convention —
database calls only in private ``asyncio.run`` loops while no ``TestClient`` is alive.
"""

from __future__ import annotations

import asyncio
import uuid
from datetime import UTC, datetime, timedelta
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from typing import Any

import pytest
from conftest import needs_db
from openpyxl import load_workbook

from app.core.db import db
from app.core.security import create_access_token, hash_password
from app.services import ministry_dashboard as register, ministry_export as exports
from prisma import Json

ROUTES = Path(__file__).resolve().parents[1] / "app" / "api" / "routes"


# --------------------------------------------------------------------------------------
# Pure
# --------------------------------------------------------------------------------------


def test_every_attention_word_but_setup_is_a_query() -> None:
    assert register.ATTENTION_KEYS == ("awaitingApproval", "returned", "overdue", "setupIncomplete")
    assert register.attention_clause("returned") == {"status": "NEEDS_REVISION"}
    assert register.attention_clause("awaitingApproval") == {"status": "SUBMITTED"}
    overdue = register.attention_clause("overdue")
    assert set(overdue["status"]["in"]) == set(register.OPEN_STATUSES)
    assert register.attention_clause("setupIncomplete") is None
    assert register.attention_clause("nonsense") is None


def test_the_identity_columns_include_both_numbers_on_the_roster() -> None:
    masked = exports.identity_columns()
    assert "participant.aadhaarNumber" in masked
    assert "participant.artisanCardNo" in masked


def _entry(entity: str, data: dict[str, Any], ordinal: int = 0) -> Any:
    return SimpleNamespace(
        id=f"e-{entity}-{ordinal}", entityKey=entity, stageKey="X", ordinal=ordinal, data=data
    )


def test_the_workbook_masks_identity_numbers_and_withholds_custom_questions() -> None:
    record = SimpleNamespace(id="w1", title="Ajrakh", state="Gujarat", status="IN_PROGRESS")
    entries = [
        _entry(
            "participant",
            {"name": "Sita", "artisanCardNo": "PC12345678", "aadhaarNumber": "234567890124"},
        ),
        _entry("_custom", {"q1": "anything"}),
    ]
    payload, rows = exports.build_workbook(
        [(record, entries)],
        filters={"state": "Gujarat"},
        scope_label="Every workshop on the platform.",
        workshops_total=1,
        generated_at=datetime(2026, 10, 17, tzinfo=UTC),
    )
    assert rows == 1
    wb = load_workbook(BytesIO(payload))
    readme = [tuple(row) for row in wb.worksheets[0].iter_rows(values_only=True)]
    assert ("Identity numbers", "Masked to the last four characters.") in readme
    assert ("Designers' own questions", "1 rows not included.") in readme
    text = "\n".join(
        str(cell)
        for sheet in wb.worksheets[1:]
        for row in sheet.iter_rows(values_only=True)
        for cell in row
        if cell
    )
    assert "PC12345678" not in text and "234567890124" not in text
    assert "5678" in text and "0124" in text
    assert "Sita" in text


def test_the_workbook_built_in_batches_matches_the_one_built_at_once() -> None:
    """The job reads and flattens one batch of workshops at a time; the file must not change."""
    first = SimpleNamespace(id="w1", title="Ajrakh", state="Gujarat", status="IN_PROGRESS")
    second = SimpleNamespace(id="w2", title="Ikat", state="Odisha", status="IN_PROGRESS")
    pairs = [
        (first, [_entry("participant", {"name": "Sita", "aadhaarNumber": "234567890124"})]),
        (second, [_entry("participant", {"name": "Gita"}), _entry("_custom", {"q": "x"})]),
    ]
    at = datetime(2026, 10, 17, tzinfo=UTC)
    whole = exports.build_workbook(
        pairs, filters={}, scope_label="All.", workshops_total=2, generated_at=at
    )
    collector = exports.WorkbookCollector()
    collector.add(pairs[:1])
    collector.add(pairs[1:])
    batched = collector.workbook(filters={}, scope_label="All.", workshops_total=2, generated_at=at)
    assert batched[1] == whole[1] == 2

    def cells(payload: bytes) -> list[tuple[Any, ...]]:
        wb = load_workbook(BytesIO(payload))
        return [tuple(row) for sheet in wb.worksheets for row in sheet.iter_rows(values_only=True)]

    assert cells(batched[0]) == cells(whole[0])


def test_the_job_list_never_reads_the_stored_file_and_the_build_is_off_the_loop() -> None:
    """The list is polled while a job runs; each stored workbook is megabytes."""
    assert '"content"' not in exports._LIST_SQL
    service = (ROUTES.parents[1] / "services" / "ministry_export.py").read_text(encoding="utf-8")
    body = service[service.index("async def run_export(") : service.index("async def recover_and_purge(")]
    assert "asyncio.to_thread(" in body
    assert "EXPORT_READ_BATCH" in body
    assert "ministrydataexport.update(" not in body.split("ministrydataexport.update_many(")[-1]
    route = (ROUTES / "ministry_stage_data_exports.py").read_text(encoding="utf-8")
    listing = route[route.index("async def list_stage_data_exports(") :]
    listing = listing[: listing.index("@router.get", 1)]
    assert "find_many" not in listing and "exports.list_jobs(" in listing


def test_a_raw_job_row_reads_like_a_model_row() -> None:
    row = {
        "id": "e1",
        "status": "READY",
        "filters": '{"state": "Gujarat"}',
        "workshopCount": 1,
        "rowCount": 2,
        "fileName": "f.xlsx",
        "fileSizeBytes": 10,
        "error": None,
        "createdAt": "2026-10-17T10:00:00+00:00",
        "finishedAt": None,
        "expiresAt": None,
    }
    payload = exports.job_payload(row)
    assert payload["filters"] == {"state": "Gujarat"}
    assert payload["createdAt"] == "2026-10-17T10:00:00+00:00"
    assert payload["finishedAt"] is None
    assert "content" not in payload


def test_the_export_routes_are_their_own_module_and_the_register_stays_read_only() -> None:
    register_source = (ROUTES / "ministry_dashboard.py").read_text(encoding="utf-8")
    assert "@router.post" not in register_source
    export_source = (ROUTES / "ministry_stage_data_exports.py").read_text(encoding="utf-8")
    assert export_source.count("Depends(require_ministry_dashboard_reader)") == 3


# --------------------------------------------------------------------------------------
# Against Postgres
# --------------------------------------------------------------------------------------


@pytest.fixture(scope="module")
def world() -> Any:
    from fastapi.testclient import TestClient

    from app.main import app

    stamp = uuid.uuid4().hex[:8]
    facts: dict[str, Any] = {"stamp": stamp}
    state = "Mizoram"

    async def seed() -> None:
        await db.connect()
        try:
            for slug, role in (
                ("ministry", "MINISTRY_ADMIN"),
                ("director", "ASSISTANT_DIRECTOR"),
                ("designer", "DESIGNER"),
            ):
                email = f"c14-{slug}-{stamp}@example.org"
                facts[slug] = await db.user.create(
                    data={
                        "email": email,
                        "name": f"C14 {slug} {stamp}",
                        "role": role,
                        "passwordHash": hash_password("c14-password-123"),
                    }
                )
                await db.accessroster.create(
                    data={
                        "email": email,
                        "status": "ACTIVE",
                        "admitRole": role,
                        "joinedAt": datetime.now(UTC),
                        "notes": "Seeded by tests/test_ministry_dashboard_attention_export.py.",
                    }
                )
            creator = facts["ministry"].id
            past = datetime.now(UTC) - timedelta(days=30)

            async def workshop(name: str, **data: Any) -> Any:
                return await db.designworkshop.create(
                    data={
                        "title": f"C14 {name} {stamp}",
                        "templateId": "DCH_STANDARD",
                        "createdById": creator,
                        **data,
                    }
                )

            facts["returned"] = await workshop("returned", status="NEEDS_REVISION", state=state)
            facts["overdue"] = await workshop(
                "overdue", status="IN_PROGRESS", endDate=past, state=state
            )
            facts["handed"] = await workshop("handed", status="SUBMITTED", endDate=past)
            await db.dwstageentry.create(
                data={
                    "designWorkshopId": facts["overdue"].id,
                    "stageKey": "WORKSHOP_PLAN_PARTICIPANTS_OPENING",
                    "entityKey": "participant",
                    "data": Json({"name": "Lalthan", "artisanCardNo": "MZ99887766"}),
                }
            )
            designer = facts["designer"]
            await db.sanctionorder.create(
                data={
                    "sanctionOrderNo": f"SO-C14-{stamp}",
                    "sanctionOrderKey": f"SOC14{stamp}".upper(),
                    "sanctionOrderDate": datetime(2026, 10, 1, tzinfo=UTC),
                    "sanctionAmount": "125000.50",
                    "designerUserId": designer.id,
                    "designerEmail": designer.email,
                    "designWorkshopId": facts["returned"].id,
                    "createdById": creator,
                }
            )
        finally:
            await db.disconnect()

    asyncio.run(seed())

    def headers(slug: str) -> dict[str, str]:
        return {"Authorization": f"Bearer {create_access_token(subject=facts[slug].id)}"}

    listing = "/api/ministry-dashboard/design-workshops"
    with TestClient(app) as client:
        facts["summary"] = client.get(
            "/api/ministry-dashboard/summary", headers=headers("ministry")
        )
        facts["list_returned"] = client.get(
            listing, params={"search": stamp, "attention": "returned"}, headers=headers("ministry")
        )
        facts["list_overdue"] = client.get(
            listing, params={"search": stamp, "attention": "overdue"}, headers=headers("ministry")
        )
        facts["list_state"] = client.get(
            listing, params={"search": stamp, "state": state}, headers=headers("ministry")
        )
        facts["list_setup"] = client.get(
            listing,
            params={"search": stamp, "attention": "setupIncomplete"},
            headers=headers("ministry"),
        )
        exports_url = "/api/ministry-dashboard/stage-data-exports"
        facts["designer_request"] = client.post(exports_url, json={}, headers=headers("designer"))
        facts["requested"] = client.post(
            exports_url, json={"search": stamp, "attention": "overdue"}, headers=headers("ministry")
        )

    async def build() -> None:
        await db.connect()
        try:
            await exports.process_next_ministry_exports(worker_id="test", limit=50)
        finally:
            await db.disconnect()

    asyncio.run(build())

    with TestClient(app) as client:
        export_id = (
            facts["requested"].json().get("id") if facts["requested"].status_code == 202 else ""
        )
        facts["listed"] = client.get(exports_url, headers=headers("ministry"))
        facts["download"] = client.get(
            f"{exports_url}/{export_id}/download", headers=headers("ministry")
        )
        facts["someone_else"] = client.get(
            f"{exports_url}/{export_id}/download", headers=headers("director")
        )
    yield facts


def _titles(response: Any) -> set[str]:
    return {row["title"] for row in response.json()["items"]}


@needs_db
def test_the_summary_carries_the_cards_the_states_and_the_rupees(world: Any) -> None:
    assert world["summary"].status_code == 200, world["summary"].text
    half = world["summary"].json()["designWorkshops"]
    assert set(half["attention"]) == set(register.ATTENTION_KEYS)
    assert half["attention"]["returned"] >= 1
    assert half["attention"]["overdue"] >= 1
    assert half["attention"]["awaitingApproval"] >= 1
    assert any(row["state"] == "Mizoram" and row["count"] >= 2 for row in half["byState"])
    sanctioned = half["sanctioned"]
    assert isinstance(sanctioned["amount"], str)
    assert Decimal(sanctioned["amount"]) >= Decimal("125000.50")


@needs_db
def test_each_card_opens_the_list_it_counted(world: Any) -> None:
    stamp = world["stamp"]
    assert _titles(world["list_returned"]) == {f"C14 returned {stamp}"}
    assert _titles(world["list_overdue"]) == {f"C14 overdue {stamp}"}
    assert _titles(world["list_state"]) == {f"C14 returned {stamp}", f"C14 overdue {stamp}"}
    # Neither open workshop has filled Workshop setup; the handed-in one is not "open".
    assert _titles(world["list_setup"]) == {f"C14 returned {stamp}", f"C14 overdue {stamp}"}


@needs_db
def test_a_designer_cannot_ask_for_the_export(world: Any) -> None:
    assert world["designer_request"].status_code == 403


@needs_db
def test_the_export_is_built_in_the_background_and_masked(world: Any) -> None:
    assert world["requested"].status_code == 202, world["requested"].text
    assert world["requested"].json()["status"] == exports.QUEUED
    item = next(
        row
        for row in world["listed"].json()["items"]
        if row["id"] == world["requested"].json()["id"]
    )
    assert item["status"] == exports.READY, item
    assert item["workshopCount"] == 1
    # Read without the file (a raw select): the dates and the filters still come through.
    assert item["createdAt"] and item["finishedAt"] and item["expiresAt"]
    assert item["filters"].get("attention") == "overdue"
    assert world["download"].status_code == 200, world["download"].text
    wb = load_workbook(BytesIO(world["download"].content))
    text = "\n".join(
        str(cell)
        for sheet in wb.worksheets
        for row in sheet.iter_rows(values_only=True)
        for cell in row
        if cell
    )
    assert "Lalthan" in text
    assert "MZ99887766" not in text and "7766" in text


@needs_db
def test_another_reader_cannot_download_somebody_elses_export(world: Any) -> None:
    assert world["someone_else"].status_code == 404
