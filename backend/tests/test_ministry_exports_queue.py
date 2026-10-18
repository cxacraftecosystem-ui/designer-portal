"""THE MINISTRY PACKS' WORK QUEUE, OVER A DATABASE.

Claiming is a compare-and-set, so of two claimers one gets the row; a build abandoned by a dead
worker comes back; the last attempt fails for good; regenerating bumps the revision and writes new
keys beside the old; a scheme with no workshops is EMPTY with no files; and a storage failure is
retried without ever leaving a READY row that points at nothing.
"""

from __future__ import annotations

import asyncio
from datetime import UTC, datetime, timedelta
from io import BytesIO
from typing import Any

import pytest
from conftest import needs_db
from ministry_exports_support import (
    FakeBucket,
    make_user,
    make_workshop,
    private,
    reset,
    settings_on,
)

import app.services.stage_definitions  # noqa: F401  - installs the registry
from app.core.db import db
from app.services import ministry_exports as packs

pytestmark = needs_db

NOW = datetime(2026, 10, 1, 1, 0, tzinfo=UTC)


async def _pack(scheme: str = "OTHER", period: str = "2026-09", **extra: Any) -> Any:
    data = {
        "period": period,
        "scheme": scheme,
        "trigger": packs.SCHEDULED,
        "status": packs.QUEUED,
        "runAfter": NOW,
        **extra,
    }
    return await db.ministryexportpack.create(data=data)


def test_of_two_claimers_exactly_one_gets_the_row() -> None:
    async def work() -> list[Any]:
        await reset()
        row = await _pack()
        return list(await asyncio.gather(packs.claim(row, "a", NOW), packs.claim(row, "b", NOW)))

    results = private(work)
    winners = [r for r in results if r is not None]
    assert len(winners) == 1
    assert winners[0].status == packs.BUILDING and winners[0].attempts == 1


def test_a_stale_build_goes_back_to_the_queue_with_a_backoff() -> None:
    async def work() -> Any:
        await reset()
        row = await _pack(
            status=packs.BUILDING, attempts=1, lockedAt=NOW - timedelta(minutes=31), lockedBy="dead"
        )
        fresh = await _pack(
            scheme="SKILL_UPGRADATION",
            status=packs.BUILDING,
            attempts=1,
            lockedAt=NOW - timedelta(minutes=5),
        )
        assert await packs.recover_stale(NOW) == 1
        return await db.ministryexportpack.find_unique(
            where={"id": row.id}
        ), await db.ministryexportpack.find_unique(where={"id": fresh.id})

    stale, fresh = private(work)
    assert stale.status == packs.QUEUED and stale.lockedBy is None
    assert stale.runAfter.astimezone(UTC) == NOW + timedelta(minutes=5)
    assert fresh.status == packs.BUILDING


def test_a_stale_build_on_its_last_attempt_fails() -> None:
    async def work() -> Any:
        await reset()
        row = await _pack(status=packs.BUILDING, attempts=3, lockedAt=NOW - timedelta(hours=1))
        await packs.recover_stale(NOW)
        return await db.ministryexportpack.find_unique(where={"id": row.id})

    row = private(work)
    assert row.status == packs.FAILED and row.error


def test_failed_after_max_attempts(monkeypatch: pytest.MonkeyPatch) -> None:
    async def broken(pack: Any, settings: Any = None) -> None:
        raise RuntimeError("renderer fell over")

    monkeypatch.setattr(packs, "build_pack", broken)

    async def work() -> list[Any]:
        await reset()
        await _pack(maxAttempts=2)
        seen = []
        moment = NOW
        for _ in range(2):
            row = await packs.process_next_pack(settings_on(), now=moment)
            seen.append(row)
            moment = moment + timedelta(hours=2)
        return seen

    first, second = private(work)
    assert first.status == packs.QUEUED and first.attempts == 1
    assert first.runAfter.astimezone(UTC) == NOW + timedelta(minutes=5)
    assert second.status == packs.FAILED and second.attempts == 2
    assert "renderer fell over" not in (
        second.error or ""
    )  # a short reason, never the exception text


def _xlsx_sheets(body: bytes) -> list[str]:
    from openpyxl import load_workbook

    return load_workbook(BytesIO(body)).sheetnames


def test_a_pack_is_built_and_regenerating_writes_new_keys(monkeypatch: pytest.MonkeyPatch) -> None:
    bucket = FakeBucket().install(monkeypatch)

    async def work() -> tuple[Any, Any]:
        await reset()
        owner = await make_user("ADMIN", "owner")
        await make_workshop(owner.id, "EXPOSURE_EXHIBITION", status="APPROVED", state="Odisha")
        await make_workshop(owner.id, "EXPOSURE_EXHIBITION", status="IN_PROGRESS", state="Assam")
        await _pack(scheme="EXPOSURE_EXHIBITION")
        first = await packs.process_next_pack(settings_on(), now=NOW)
        [again] = await packs.request_packs("2026-09", ["EXPOSURE_EXHIBITION"], owner.id)
        assert again.revision == 2 and again.status == packs.QUEUED and again.xlsxKey is None
        second = await packs.process_next_pack(settings_on())
        return first, second

    first, second = private(work)
    assert first.status == packs.READY and first.revision == 1
    assert first.counts["workshops"] >= 2
    assert "/r1/" in first.xlsxKey and "/r2/" in second.xlsxKey
    assert second.status == packs.READY and second.trigger == packs.ON_DEMAND
    assert {first.xlsxKey, first.pdfKey, second.xlsxKey, second.pdfKey} <= set(bucket.objects)
    assert bucket.objects[first.pdfKey].startswith(b"%PDF")
    assert _xlsx_sheets(bucket.objects[second.xlsxKey])[0] == "Summary"
    assert first.xlsxBytes == len(bucket.objects[first.xlsxKey]) and len(first.xlsxSha256) == 64


def test_a_scheme_with_no_workshops_is_empty_and_has_no_files(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    bucket = FakeBucket().install(monkeypatch)

    async def work() -> Any:
        await reset()
        await db.designworkshop.update_many(
            where={"workshopKind": "CLUSTER_DEVELOPMENT", "deletedAt": None},
            data={"deletedAt": NOW},
        )
        await _pack(scheme="CLUSTER_DEVELOPMENT")
        return await packs.process_next_pack(settings_on(), now=NOW)

    row = private(work)
    assert row.status == packs.EMPTY
    assert row.xlsxKey is None and row.pdfKey is None
    assert row.counts["workshops"] == 0
    assert not bucket.objects


def test_a_storage_failure_is_retried_and_never_leaves_a_ready_row_without_files(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    bucket = FakeBucket().install(monkeypatch)
    bucket.fail_puts = 1

    async def work() -> tuple[Any, Any]:
        await reset()
        owner = await make_user("ADMIN", "owner")
        await make_workshop(owner.id, "DESIGN_INTERVENTION", status="SUBMITTED")
        await _pack(scheme="DESIGN_INTERVENTION")
        failed = await packs.process_next_pack(settings_on(), now=NOW)
        retried = await packs.process_next_pack(settings_on(), now=NOW + timedelta(minutes=6))
        return failed, retried

    failed, retried = private(work)
    assert failed.status == packs.QUEUED and failed.xlsxKey is None and failed.error
    assert retried.status == packs.READY
    assert retried.xlsxKey in bucket.objects and retried.pdfKey in bucket.objects
    assert retried.error is None


def test_unstated_selects_workshops_with_no_kind(monkeypatch: pytest.MonkeyPatch) -> None:
    FakeBucket().install(monkeypatch)

    async def work() -> tuple[int, int]:
        await reset()
        owner = await make_user("ADMIN", "owner")
        await make_workshop(owner.id, None, status="DRAFT")
        await _pack(scheme="UNSTATED")
        row = await packs.process_next_pack(settings_on(), now=NOW)
        expected = await db.designworkshop.count(where={"workshopKind": None, "deletedAt": None})
        return int(row.counts["workshops"]), expected

    got, expected = private(work)
    assert got == min(expected, settings_on().ministry_exports_max_workshops) and got >= 1


def test_the_cap_truncates_and_says_so(monkeypatch: pytest.MonkeyPatch) -> None:
    FakeBucket().install(monkeypatch)

    async def work() -> Any:
        await reset()
        owner = await make_user("ADMIN", "owner")
        for _ in range(3):
            await make_workshop(owner.id, "SKILL_UPGRADATION", status="DRAFT")
        await _pack(scheme="SKILL_UPGRADATION")
        return await packs.process_next_pack(settings_on(ministry_exports_max_workshops=2), now=NOW)

    row = private(work)
    assert row.counts["workshops"] == 2 and row.counts["truncated"] is True
