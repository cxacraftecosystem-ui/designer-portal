"""THE SANCTION-ORDER SHEET READER — the pro-forma it writes and the workbook it reads back.

NO DATABASE, AND THAT IS STRUCTURAL RATHER THAN CONVENIENT. ``sanction_orders_xlsx`` deliberately
holds no reference to ``services/sanction_orders`` — it answers "what do the cells say" and nothing
about who the people are — so every assertion here runs against real openpyxl bytes with no fixture,
no Prisma client and no network.

══ THE THREE FAILURES THIS FILE EXISTS FOR ══════════════════════════════════════════════════════

1. **A COLUMN PRINTED ON THE PRO-FORMA THAT THE READER CANNOT FIND.** An office fills it in, the
   app stores nothing from it, and the upload answers 200. There is no symptom anywhere. The module
   asserts the invariant at IMPORT time; :func:`test_every_pro_forma_heading_is_one_this_reader_finds`
   asserts it again here so that the failure names the column rather than arriving as a collection
   error. **This has already fired once, on the first run**: ``norm_header("Designer Name(s)")`` is
   ``"designer name s"`` — the parentheses fold to spaces and the "s" stands alone — so the blank
   pro-forma shipped with a designer column its own reader could not match.

2. **THE MONEY BECOMING A FLOAT.** A cell Excel stored as a number arrives as a Python ``float``,
   and ``Decimal(0.1)`` is ``0.1000000000000000055511151231257827``. The register is
   ``NUMERIC(14,2)`` and the wire is a decimal string precisely to keep binary floats out of a
   ministry budget; this reader is the fourth layer of that rule and the easiest one to get wrong,
   because ``Decimal(value)`` compiles and looks right.

3. **THE TWO DESIGNER COLUMNS LOSING THEIR CORRESPONDENCE.** The Nth name goes with the Nth address.
   One trailing comma, and every pairing below it is wrong while every cell still looks reasonable.
   The split is where that is either preserved or silently destroyed.
"""

from __future__ import annotations

import inspect
import io
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook

from app.services import sanction_orders_xlsx as reader
from app.services.sanction_orders_xlsx import (
    MAX_DESIGNERS_PER_ROW,
    MAX_SANCTION_ROWS,
    SanctionXlsxError,
    build_sanction_pro_forma,
    fold_name,
    fold_order_no,
    parse_sanction_workbook,
    split_list_cell,
)
from app.services.xlsx_table import FormulaCell, ParseProblem, norm_header

HEADINGS = [
    "Sanction Order No.",
    "Sanction Date",
    "Sanctioned Amount",
    "Designer Name(s)",
    "Designer Email(s)",
    "Remarks",
]


def sheet_of(rows: list[tuple], *, headings: list[str] | None = None) -> bytes:
    """A workbook with the pro-forma's headings and the given rows under them."""
    wb = Workbook()
    ws = wb.active
    ws.title = reader.SHEET_ORDERS
    for index, label in enumerate(headings or HEADINGS, start=1):
        ws.cell(row=1, column=index, value=label)
    for r, values in enumerate(rows, start=2):
        for c, value in enumerate(values, start=1):
            ws.cell(row=r, column=c, value=value)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def problems_on(parsed, row: int, severity: str | None = None) -> list[ParseProblem]:
    return [
        problem
        for problem in parsed.problems
        if problem.row == row and (severity is None or problem.severity == severity)
    ]


# --------------------------------------------------------------------------------------
# 1. The pro-forma and the reader agree about every column
# --------------------------------------------------------------------------------------


def test_every_pro_forma_heading_is_one_this_reader_finds() -> None:
    """THE INVARIANT THAT CAUGHT A REAL DEFECT ON THE FIRST RUN OF THIS MODULE.

    A heading printed on the workbook this product HANDS OUT, which no alias matches, is a column an
    office fills in and the app silently ignores — under a 200, with nothing on any screen. The
    module asserts this at import time so the process refuses to start; this asserts it again by
    NAME so that a failure says which column rather than arriving as a collection error nobody can
    read.

    ``norm_header`` is applied on both sides because that is what the matcher applies: it strips
    punctuation TO SPACES, so "Designer Name(s)" is ``"designer name s"``. Comparing the raw labels
    would pass while the reader found nothing.
    """
    for role, label in reader._PRO_FORMA_COLUMNS:
        assert norm_header(label) in reader._COLUMN_ALIASES[role], (role, label, norm_header(label))


def test_the_pro_forma_has_no_row_anybody_could_upload() -> None:
    """Headings and nothing under them.

    A seeded example row is a row somebody uploads, and an imported sanction order mints accounts,
    issues credentials and opens a workshop — none of which this product can delete. The worked
    example lives on the instructions sheet, where no cell normalises to a column alias and
    ``pick_sheet`` therefore cannot choose it as the data sheet.
    """
    wb = load_workbook(io.BytesIO(build_sanction_pro_forma()))
    assert wb.sheetnames == [reader.SHEET_ORDERS, reader.SHEET_HELP]
    orders = wb[reader.SHEET_ORDERS]
    assert [cell.value for cell in orders[1]] == HEADINGS
    assert orders.max_row == 1, [
        [cell.value for cell in row] for row in orders.iter_rows(min_row=2)
    ]


def test_the_instructions_sheet_is_not_even_a_candidate_for_the_data_sheet() -> None:
    """No cell on it normalises to a column alias, so ``pick_sheet`` cannot pick it.

    That is the property that lets the worked example be a worked example rather than a trap: every
    line is ONE cell of piped text, so there is no row on that sheet that looks like a header.
    """
    wb = load_workbook(io.BytesIO(build_sanction_pro_forma()))
    every_alias = {alias for aliases in reader._COLUMN_ALIASES.values() for alias in aliases}
    for row in wb[reader.SHEET_HELP].iter_rows():
        for cell in row:
            assert norm_header(cell.value) not in every_alias, cell.value


def test_the_pro_forma_round_trips_through_its_own_reader() -> None:
    """Write the blank workbook, type into it, read it back. The whole feature in one assertion."""
    wb = load_workbook(io.BytesIO(build_sanction_pro_forma()))
    ws = wb[reader.SHEET_ORDERS]
    ws.append(["SO/2026/42", date(2026, 4, 3), 450000, "Ramesh Kumar", "r.kumar@gmail.com", "ok"])
    buffer = io.BytesIO()
    wb.save(buffer)

    parsed = parse_sanction_workbook(buffer.getvalue(), filename="orders.xlsx")
    assert parsed.sheet == reader.SHEET_ORDERS
    assert parsed.sourceFilename == "orders.xlsx"
    assert len(parsed.rows) == 1
    row = parsed.rows[0]
    assert row.sanctionOrderNo == "SO/2026/42"
    assert row.sanctionOrderDate == date(2026, 4, 3)
    # SPELT AS A STRING AND IN PAISE. ``Decimal(450000)`` is what ruff would rather see and it is
    # the example the next reader copies — into a float. The rule this module is written against
    # is that a money value is never built from a number; the two-decimal form also says what a
    # NUMERIC(14,2) column actually holds. Decimal equality ignores the exponent, so this is the
    # same assertion.
    assert row.sanctionAmount == Decimal("450000.00")
    assert row.designerNames == ["Ramesh Kumar"]
    assert row.designerEmails == ["r.kumar@gmail.com"]
    assert row.notes == "ok"
    # A REAL DATE CELL PRODUCES NO WARNING. The day-first notice is for TEXT dates, where both
    # readings are possible; Excel has already decided for a real one and second-guessing it is the
    # corruption ``read_plan_date`` exists to prevent.
    assert problems_on(parsed, 2) == []


# --------------------------------------------------------------------------------------
# 2. The bounds mirror the schema the values are written into
# --------------------------------------------------------------------------------------


def test_no_clip_cap_exceeds_the_field_it_is_written_into() -> None:
    """Every ``_TEXT_CAPS`` number is its ``SanctionOrderCreate`` field's, and never looser.

    A cap one character LOOSER than its field is the expensive direction: the row parses, the officer
    agrees to it on the confirmation screen, and the 422 arrives from pydantic at the moment the
    order is being written — by which time they have already made a decision about it.

    Read off the model's own ``max_length`` rather than restated, so the pair cannot be edited apart.
    This is ``test_no_parser_clip_cap_exceeds_the_field_it_is_promoted_into``'s shape with the
    registry swapped for the schema, because nothing on a sanction sheet is promoted into a stage.
    """
    from app.schemas.sanction_orders import SanctionOrderCreate

    for role, cap in reader._TEXT_CAPS.items():
        field_name = reader._CAP_MIRRORS[role]
        field = SanctionOrderCreate.model_fields[field_name]
        bounds = [
            meta.max_length
            for meta in field.metadata
            if getattr(meta, "max_length", None) is not None
        ]
        assert bounds, field_name
        assert cap <= bounds[0], (role, field_name, cap, bounds[0])


def test_the_in_file_fold_is_the_registers_own_fold() -> None:
    """``fold_order_no`` and ``normalise_sanction_order_no`` agree over every house spelling.

    The parser spells the fold itself rather than importing the service, so that this whole module
    stays free of Prisma. That duplication is only safe while the two agree: if they diverged, two
    rows of one sheet could pass the in-file duplicate check and then collide on
    ``@@unique(sanctionOrderKey)`` — inside a transaction that has already minted accounts and
    admitted people to the platform.
    """
    from app.services.sanction_orders import normalise_sanction_order_no

    for spelling in (
        "SO/2026/42",
        "SO-2026-42",
        "so 2026 42",
        "  so/2026/42  ",
        "So.2026.42",
        "",
        "///",
        "2026",
    ):
        assert fold_order_no(spelling) == normalise_sanction_order_no(spelling), spelling


# --------------------------------------------------------------------------------------
# 3. The money is never a float
# --------------------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("cell", "expected"),
    [
        (450000, Decimal("450000.00")),
        (450000.10, Decimal("450000.1")),
        ("450000.00", Decimal("450000.00")),
        ("6,25,000/-", Decimal("625000.00")),
        ("Rs. 1,00,000", Decimal("100000.00")),
        ("₹ 45,000.50", Decimal("45000.50")),
        ("INR 45000", Decimal("45000.00")),
    ],
)
def test_the_amount_survives_every_spelling_an_office_uses(cell, expected) -> None:
    """₹, "Rs.", "INR", thousands commas and the Indian "/-" are noise around the number."""
    problems: list[ParseProblem] = []
    assert reader.read_amount(cell, sheet="s", row=2, problems=problems) == expected
    assert [p for p in problems if p.severity == "error"] == []


def test_a_numeric_cell_never_becomes_a_binary_float() -> None:
    """``0.1`` off a spreadsheet is ``Decimal("0.1")`` and not ``Decimal(0.1)``.

    THE LINE THIS PROTECTS IS ONE CHARACTER WIDE. ``Decimal(raw)`` on a float compiles, runs and
    produces ``0.1000000000000000055511151231257827``; ``Decimal(str(raw))`` produces ``0.1``. The
    register is ``NUMERIC(14,2)`` and the wire is a decimal string precisely so that a ministry
    budget is never a binary float, and this is the layer where it would re-enter.
    """
    problems: list[ParseProblem] = []
    got = reader.read_amount(0.1, sheet="s", row=2, problems=problems)
    assert got == Decimal("0.1")
    assert str(got) == "0.1"
    # THE POINT OF THIS LINE IS THE FLOAT, so it is spelled as one and the rule that
    # normally forbids it is suppressed by name. ``Decimal(0.1)`` is
    # 0.1000000000000000055511151231257827 and is exactly the value this reader must never
    # produce; asserting inequality against a STRING literal here would assert nothing.
    assert got != Decimal(0.1)  # noqa: RUF032


def test_more_than_two_decimal_places_is_rounded_and_said_rather_than_refused() -> None:
    """A third decimal place is a spreadsheet's rounding artefact far more often than an intention.

    Refusing the row would lose a ministry order over a digit nobody typed; rounding it silently
    would change a figure on a financial instrument. So: quantised to paise, and reported.
    """
    problems: list[ParseProblem] = []
    assert reader.read_amount("100.005", sheet="s", row=2, problems=problems) == Decimal("100.00")
    assert len(problems) == 1
    assert problems[0].severity == "warning"
    assert "two decimal places" in problems[0].reason


@pytest.mark.parametrize("cell", ["abc", "45000 to 50000", "forty five thousand"])
def test_an_amount_with_no_single_figure_in_it_is_refused_and_not_guessed(cell) -> None:
    """Which figure the ministry sanctioned is not a decision a parser takes."""
    problems: list[ParseProblem] = []
    assert reader.read_amount(cell, sheet="s", row=2, problems=problems) is None
    assert problems and problems[0].severity == "error"


def test_zero_and_negative_amounts_are_refused() -> None:
    """``SanctionOrderCreate.sanctionAmount`` is ``gt=0``; this is that bound met earlier."""
    for cell in (0, "0.00", -5):
        problems: list[ParseProblem] = []
        assert reader.read_amount(cell, sheet="s", row=2, problems=problems) is None
        assert problems and problems[0].severity == "error"


def test_a_formula_with_no_cached_result_is_an_error_and_not_an_empty_cell() -> None:
    """"Row 12 is blank" about a row that visibly reads ``=B12*2`` is the worst thing to say.

    A workbook written by a script has formulas and no cached values. The distinction matters more
    on a money column than anywhere else: an amount that read as blank would refuse the order, and
    an officer looking at a filled-in cell would have nothing to go on.
    """
    problems: list[ParseProblem] = []
    assert reader.read_amount(FormulaCell("=B2*2"), sheet="s", row=2, problems=problems) is None
    assert problems[0].severity == "error"
    assert "formula" in problems[0].reason


# --------------------------------------------------------------------------------------
# 4. The two columns that correspond positionally
# --------------------------------------------------------------------------------------


def test_a_comma_is_the_only_separator() -> None:
    problems: list[ParseProblem] = []
    got = split_list_cell(
        "Ramesh Kumar, Kavita Rao ,Anil Shah",
        sheet="s",
        row=2,
        column_label="Designer Name(s)",
        problems=problems,
    )
    assert got == ["Ramesh Kumar", "Kavita Rao", "Anil Shah"]
    assert problems == []


@pytest.mark.parametrize("cell", ["Ramesh\nKavita", "Ramesh;Kavita", "Ramesh\r\nKavita"])
def test_a_line_break_or_a_semicolon_is_reported_and_never_split_on(cell) -> None:
    """ONE ENTRY, and a warning naming what was found.

    Splitting on a newline silently would make two people out of one whose name happens to wrap, or
    one out of two. Neither has a symptom. So the cell is read as it stands — with the break folded
    to a space, because a stored name with a line break in it is its own small disaster — and the
    officer is told what their cell actually contains.
    """
    problems: list[ParseProblem] = []
    got = split_list_cell(
        cell, sheet="s", row=2, column_label="Designer Name(s)", problems=problems
    )
    assert got == ["Ramesh Kavita"]
    assert len(problems) == 1
    assert problems[0].severity == "warning"
    assert "line break or a semicolon" in problems[0].reason


def test_an_empty_segment_is_dropped_and_counted_before_anything_compares_the_counts() -> None:
    """``"Ramesh, , Kavita"`` is two names, not three, and the officer is told why.

    THE ORDERING IS THE WHOLE POINT and is what this test is really about. Left in, a trailing comma
    in one cell and not the other manufactures a COUNT MISMATCH out of a typo — and a count mismatch
    is a question put to a human. Dropping these first is what keeps the confirmation step for rows
    that genuinely need a decision.
    """
    problems: list[ParseProblem] = []
    assert split_list_cell(
        "Ramesh, , Kavita", sheet="s", row=2, column_label="Designer Name(s)", problems=problems
    ) == ["Ramesh", "Kavita"]
    assert problems[0].severity == "warning"
    assert "empty" in problems[0].reason

    trailing: list[ParseProblem] = []
    assert split_list_cell(
        "a@x.com,b@x.com,", sheet="s", row=2, column_label="Designer Email(s)", problems=trailing
    ) == ["a@x.com", "b@x.com"]
    assert trailing and trailing[0].severity == "warning"


def test_whitespace_of_every_invisible_kind_is_folded_silently() -> None:
    """A non-breaking space, a zero-width space, a full-width space: folded, and NOT reported.

    These are mechanical and there is nothing for an officer to decide about them — reporting each
    one would be the "warns about everything" failure. What IS reported is anything that could have
    been a separator.
    """
    problems: list[ParseProblem] = []
    got = split_list_cell(
        "Ramesh \u200b Kumar ,\u3000Kavita Rao",
        sheet="s",
        row=2,
        column_label="Designer Name(s)",
        problems=problems,
    )
    assert got == ["Ramesh Kumar", "Kavita Rao"]
    assert problems == []


def test_a_cell_with_a_whole_column_pasted_into_it_is_cut_and_said() -> None:
    """A guard against a malformed file, never a silent truncation."""
    problems: list[ParseProblem] = []
    got = split_list_cell(
        ",".join(f"n{i}" for i in range(MAX_DESIGNERS_PER_ROW + 5)),
        sheet="s",
        row=2,
        column_label="Designer Name(s)",
        problems=problems,
    )
    assert len(got) == MAX_DESIGNERS_PER_ROW
    assert problems and problems[0].severity == "error"


def test_the_two_columns_keep_their_positions_through_the_reader() -> None:
    """The Nth name comes back with the Nth address, in the order they were typed."""
    parsed = parse_sanction_workbook(
        sheet_of(
            [
                (
                    "SO/1",
                    "03/04/2026",
                    "1000",
                    "Kavita Rao, Anil Shah, Ramesh Kumar",
                    "k@x.com, a@x.com, r@x.com",
                    None,
                )
            ]
        )
    )
    row = parsed.rows[0]
    assert row.designerNames == ["Kavita Rao", "Anil Shah", "Ramesh Kumar"]
    assert row.designerEmails == ["k@x.com", "a@x.com", "r@x.com"]


def test_fold_name_is_for_comparison_and_never_for_storage() -> None:
    """Two spellings of one name fold together; what the reader STORES is what was typed."""
    assert fold_name("RAMESH  KUMAR") == fold_name("Ramesh Kumar")
    assert fold_name("Ramesh Kumar") == fold_name("ramesh kumar")
    parsed = parse_sanction_workbook(
        sheet_of([("SO/1", "03/04/2026", "1000", "RAMESH  KUMAR", "r@x.com", None)])
    )
    assert parsed.rows[0].designerNames == ["RAMESH KUMAR"]


# --------------------------------------------------------------------------------------
# 5. Rows, and the ones that never become rows
# --------------------------------------------------------------------------------------


def test_a_row_with_no_order_number_is_reported_and_not_dropped_in_silence() -> None:
    """The order number is what the register is keyed on; without it there is nothing to file."""
    parsed = parse_sanction_workbook(
        sheet_of(
            [
                ("SO/1", "03/04/2026", "1000", "A", "a@x.com", None),
                (None, "04/04/2026", "2000", "B", "b@x.com", None),
            ]
        )
    )
    assert [row.sanctionOrderNo for row in parsed.rows] == ["SO/1"]
    assert problems_on(parsed, 3, "error"), parsed.problems


def test_two_rows_with_one_number_name_each_other() -> None:
    """The in-file duplicate, which the REGISTER cannot see because neither row is written yet.

    Both Excel rows are named because the officer has to look at both to decide which is right — the
    shape ``_drop_in_file_duplicates`` uses on the artisan import. The fold is the register's own, so
    ``SO/2026/42`` and ``SO-2026-42`` collide here exactly as they would on the unique index.
    """
    parsed = parse_sanction_workbook(
        sheet_of(
            [
                ("SO/2026/42", "03/04/2026", "1000", "A", "a@x.com", None),
                ("SO-2026-42", "04/04/2026", "2000", "B", "b@x.com", None),
            ]
        )
    )
    assert len(parsed.rows) == 1
    problem = problems_on(parsed, 3, "error")[0]
    assert "row 2 of this sheet" in problem.reason


def test_a_blank_spacer_row_is_the_one_silent_skip() -> None:
    """Ministry sheets are full of them; a problem per blank line buries the two that matter."""
    parsed = parse_sanction_workbook(
        sheet_of(
            [
                ("SO/1", "03/04/2026", "1000", "A", "a@x.com", None),
                (None, None, None, None, None, None),
                ("SO/2", "04/04/2026", "2000", "B", "b@x.com", None),
            ]
        )
    )
    assert [row.sanctionOrderNo for row in parsed.rows] == ["SO/1", "SO/2"]
    assert problems_on(parsed, 3) == []


def test_a_row_whose_date_or_amount_is_unreadable_is_still_returned() -> None:
    """**IT IS RETURNED SO THE COUNTS CAN ADD UP.**

    ``rowsRead = recorded + skipped + refused`` is the arithmetic the report is checked by, and a row
    this function swallowed would be missing from all four. The reconciliation is what refuses it,
    with a sentence about the INSTRUMENT ("the date is one of the three facts it consists of") to go
    beside the parser's sentence about the CELL.
    """
    parsed = parse_sanction_workbook(
        sheet_of([("SO/1", "not a date", "abc", "A", "a@x.com", None)])
    )
    assert len(parsed.rows) == 1
    assert parsed.rows[0].sanctionOrderDate is None
    assert parsed.rows[0].sanctionAmount is None
    assert len(problems_on(parsed, 2, "error")) == 2


def test_a_sheet_beyond_the_row_ceiling_is_cut_and_says_where() -> None:
    """200 and not the annual plan's 600 — each order opens a workshop and may create an account."""
    rows = [(f"SO/{i}", "03/04/2026", "1000", "A", f"a{i}@x.com", None) for i in range(MAX_SANCTION_ROWS + 3)]
    parsed = parse_sanction_workbook(sheet_of(rows))
    assert len(parsed.rows) == MAX_SANCTION_ROWS
    ceiling = [p for p in parsed.problems if "more than" in p.reason and p.severity == "error"]
    assert ceiling, parsed.problems
    assert str(MAX_SANCTION_ROWS) in ceiling[0].reason


def test_an_unrecognised_heading_is_named_and_the_padding_is_not() -> None:
    """ONE warning naming the real headings, and no warning at all for openpyxl's ``None`` padding.

    Without the text-only filter every ordinary upload warns ", , , , ," and teaches officers to
    ignore the one panel that tells them a column of real data was skipped.
    """
    parsed = parse_sanction_workbook(
        sheet_of(
            [("SO/1", "03/04/2026", "1000", "A", "a@x.com", None, None)],
            headings=[*HEADINGS, "District Office"],
        )
    )
    warnings = [p for p in parsed.problems if p.severity == "warning" and p.row == 1]
    assert len(warnings) == 1
    assert "District Office" in warnings[0].reason
    assert ", ," not in warnings[0].reason


def test_a_workbook_with_no_order_number_column_is_refused_whole() -> None:
    """:class:`SanctionXlsxError` is the FILE, never one row — the distinction ``ParseProblem`` is for."""
    with pytest.raises(SanctionXlsxError) as caught:
        parse_sanction_workbook(sheet_of([("x",)], headings=["Something Else"]))
    assert "Sanction Order No." in str(caught.value)


def test_an_empty_upload_is_refused_with_a_sentence_and_not_a_stack_trace() -> None:
    with pytest.raises(SanctionXlsxError):
        parse_sanction_workbook(b"")


def test_a_sheet_with_the_headings_and_no_rows_is_refused_whole() -> None:
    with pytest.raises(SanctionXlsxError) as caught:
        parse_sanction_workbook(sheet_of([]))
    assert "no rows under it" in str(caught.value)


# --------------------------------------------------------------------------------------
# 6. The date rule is the annual plan's, not a second copy of it
# --------------------------------------------------------------------------------------


def test_a_text_date_is_read_day_first_and_says_so_when_both_readings_are_possible() -> None:
    """``03/04/2026`` is the third of April in every office that will send this file.

    Both readings are valid dates, neither raises, and the row imports either way — so the one case
    where this could be silently wrong is the one case the officer is shown. The warning names the
    SANCTION Date and not the Start Date, which is the whole reason ``read_plan_date`` takes a
    ``column_label``: an officer told about a column that is not on their sheet has been given a
    puzzle rather than a warning.
    """
    parsed = parse_sanction_workbook(
        sheet_of([("SO/1", "03/04/2026", "1000", "A", "a@x.com", None)])
    )
    assert parsed.rows[0].sanctionOrderDate == date(2026, 4, 3)
    warnings = problems_on(parsed, 2, "warning")
    assert warnings and "Sanction Date" in warnings[0].reason
    assert "Start Date" not in warnings[0].reason


def test_the_date_rule_is_imported_rather_than_copied() -> None:
    """One spelling of the day-first rule, shared with the annual plan.

    A second copy would be a second place for it to drift, and a drift here is a workshop scheduled
    a month out with no error anywhere in the system. Asserted by identity, so a later "inline it for
    clarity" has to delete this line.
    """
    from app.services import annual_plan_xlsx

    assert reader.read_plan_date is annual_plan_xlsx.read_plan_date


def test_this_module_never_imports_the_sanction_service() -> None:
    """The parser answers what the CELLS say and nothing about who the people are.

    That separation is what lets every test in this file run with no Prisma client, no fixture and no
    database — and it is what makes the reconciliation's own tests able to stub the register. An
    import of ``services.sanction_orders`` here would pull the database, the credential links and the
    workshop service into a module whose whole job is openpyxl.
    """
    source = inspect.getsource(reader)
    assert "from app.services.sanction_orders import" not in source
    assert "from app.services import sanction_orders" not in source
