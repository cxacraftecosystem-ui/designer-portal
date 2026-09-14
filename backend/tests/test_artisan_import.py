"""What an uploaded artisan list does to the repository, and what it must never do.

**THE DATABASE IS A FAKE AND THAT IS THE POINT, NOT A COMPROMISE.** Every assertion here is about
which WRITES the importer issues and what is in them — that no artisan is overwritten, that no
``WorkshopArtisan`` row is written, that the coordinate lands in the provenance columns and not in
the subject pin, that no twelve-digit run reaches the ledger. Those are statements about the code,
and the REAL functions are what run against these tables: ``import_artisans``,
``records.clean_data``, ``records.resolve_craft_id``, ``records.attach_location``,
``records.merge_field_provenance`` and ``records.apply_status_policy_create`` are all called for
real. A delegate narrow enough that an unexpected clause fails loudly is the same device
``test_dw_inspector_scope_gate`` uses for the predicate pair, and it means these tests RUN on a
machine with no Postgres — which is where somebody changing this code is sitting.

``design_workshops.save_stage`` IS STUBBED, and only it. It is a six-hundred-line function over
eight tables whose own behaviour has its own tests; what this file asserts is the SHAPE of the call
the importer makes into it, which is where every one of the four things a direct
``db.dwstageentry.create`` would lose is decided.
"""

from io import BytesIO
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from app.services import design_workshops, records
from app.services.artisan_import import import_artisans
from app.services.artisan_xlsx import (
    _COLUMNS,
    ArtisanDefaults,
    build_artisan_pro_forma,
    parse_artisan_workbook,
)

AADHAAR_A = "223456789018"
AADHAAR_B = "323456789012"
AADHAAR_C = "423456789019"

DEFAULTS = ArtisanDefaults(
    state="Rajasthan", district="Jaipur", place="Bagru", craftName="Block printing"
)
VENUE_FIX = {"lat": 26.8129, "lon": 75.5432, "accuracy": 12.0}

_ORDER = [key for key, _label in _COLUMNS]


# --------------------------------------------------------------------------------------
# The fake
# --------------------------------------------------------------------------------------


class _Table:
    """A delegate over a list. Narrow on purpose: an unexpected clause fails loudly.

    Only the operators the importer actually issues are implemented — equality, ``in``, and a
    top-level ``OR`` of those. A future clause shape this cannot express raises rather than quietly
    matching everything, which is the failure mode a permissive fake has.
    """

    def __init__(self, name: str, rows: list | None = None) -> None:
        self.name = name
        self.rows = list(rows or [])
        self.creates: list[dict] = []
        self.updates: list[tuple[dict, dict]] = []
        self.create_error: Exception | None = None
        self._next = 0

    # -- reading --------------------------------------------------------------------
    @staticmethod
    def _match_one(row, key, wanted) -> bool:
        value = getattr(row, key, None)
        if isinstance(wanted, dict):
            if "in" in wanted:
                return value in wanted["in"]
            raise AssertionError(f"unsupported operator {wanted!r} on {key}")
        return value == wanted

    def _matches(self, row, where: dict | None) -> bool:
        if not where:
            return True
        for key, wanted in where.items():
            if key == "OR":
                if not any(self._matches(row, clause) for clause in wanted):
                    return False
                continue
            if not self._match_one(row, key, wanted):
                return False
        return True

    async def find_many(self, where=None, include=None, order=None, take=None, skip=None):
        found = [r for r in self.rows if self._matches(r, where)]
        return found[:take] if take else found

    async def find_first(self, where=None, include=None, order=None):
        return next((r for r in self.rows if self._matches(r, where)), None)

    async def find_unique(self, where=None, include=None):
        return await self.find_first(where=where)

    # -- writing --------------------------------------------------------------------
    async def create(self, data, include=None):
        if self.create_error is not None:
            error, self.create_error = self.create_error, None
            raise error
        self.creates.append(dict(data))
        self._next += 1
        row = SimpleNamespace(id=f"{self.name}-{self._next}", **data)
        self.rows.append(row)
        return row

    async def update(self, where, data):
        self.updates.append((dict(where), dict(data)))
        return SimpleNamespace(id=where.get("id"), **data)


class _Client:
    """Every delegate the importer may touch — and an ``AttributeError`` for anything else.

    **THAT IS THE ASSERTION IN ``test_no_workshop_artisan_row_is_written``**: ``workshopartisan`` is
    deliberately absent, so a line that reached for it would raise rather than silently do the
    wrong thing on a table nobody was watching.
    """

    def __init__(self, artisans: list | None = None, crafts: list | None = None) -> None:
        self.artisan = _Table("artisan", artisans)
        # ``Block Printing`` AND NOT ``Block printing``, WHICH IS NOT A TYPO AND IS WORTH THE LINE.
        # ``records.clean_data`` title-cases ``craftName`` on EVERY write path (it is in
        # ``TITLE_CASE_FIELDS``) before ``resolve_craft_id`` looks the name up, so the spelling this
        # lookup actually sees is the title-cased one — which is also the spelling ``POST /crafts``
        # stored. A fixture holding the sheet's casing would make every create in this file refuse
        # with "no craft is called …" and the tests would be asserting the wrong thing about a real
        # behaviour.
        self.craft = _Table(
            "craft",
            crafts
            if crafts is not None
            else [SimpleNamespace(id="craft-1", name="Block Printing")],
        )
        self.location = _Table("location")
        self.dwartisanimport = _Table("import")
        self.dwstageentry = _Table("stage")
        self.designworkshop = _Table("workshop")


WORKSHOP = SimpleNamespace(
    id="cmworkshop00000000000001",
    title="Bagru block printing",
    state="Rajasthan",
    district="Jaipur",
    venue="Bagru",
    craftName="Block printing",
    createdById="creator-1",
    status="IN_PROGRESS",
)
OFFICER = SimpleNamespace(id="officer-1", name="A Director", role="MINISTRY_ADMIN")


@pytest.fixture
def world(monkeypatch: pytest.MonkeyPatch):
    """A fake client bound everywhere the importer and the record helpers reach for one."""
    import app.core.db as core_db

    client = _Client()
    saves: list = []

    async def fake_save_stage(workshop_id, spec, payload, user):
        saves.append(
            SimpleNamespace(workshop_id=workshop_id, spec=spec, payload=payload, user=user)
        )
        return {"saved": len(payload.entries), "created": len(payload.entries), "errors": {}}

    monkeypatch.setattr(core_db, "db", client)
    monkeypatch.setattr(records, "db", client)
    monkeypatch.setattr(design_workshops, "save_stage", fake_save_stage)
    return SimpleNamespace(db=client, saves=saves)


def sheet(rows: list[dict]) -> bytes:
    wb = load_workbook(BytesIO(build_artisan_pro_forma()))
    ws = wb["Artisans"]
    for r, values in enumerate(rows, start=2):
        for c, key in enumerate(_ORDER, start=1):
            if values.get(key) not in (None, ""):
                ws.cell(row=r, column=c, value=values[key])
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def artisan_row(**overrides) -> dict:
    base = {
        "name": "Ramesh Devi",
        "craftName": "Block printing",
        "place": "Bagru",
        "aadhaarNumber": AADHAAR_A,
        "dos": "Wash the cloth first",
        "donts": "Do not let the dye dry",
        "state": "Rajasthan",
        "district": "Jaipur",
    }
    base.update(overrides)
    return base


async def run_import(world, rows: list[dict], **kwargs):
    parsed = parse_artisan_workbook(sheet(rows), defaults=DEFAULTS)
    return await import_artisans(
        parsed,
        workshop=kwargs.pop("workshop", WORKSHOP),
        actor=kwargs.pop("actor", OFFICER),
        venue_fix=kwargs.pop("venue_fix", VENUE_FIX),
        source_filename=kwargs.pop("source_filename", "bagru-list.xlsx"),
    )


def stored(value):
    """A ``Json``-wrapped column as the plain value it holds."""
    return getattr(value, "data", value)


# --------------------------------------------------------------------------------------
# The happy path
# --------------------------------------------------------------------------------------


async def test_three_rows_become_three_artisans_and_three_participant_rows(world):
    report = await run_import(
        world,
        [
            artisan_row(),
            artisan_row(name="Sita Bai", aadhaarNumber=AADHAAR_B),
            artisan_row(name="Mohan Lal", aadhaarNumber=AADHAAR_C),
        ],
    )
    assert report["artisansCreated"] == 3
    assert report["rowsRefused"] == 0
    assert report["rowsRead"] == 3
    assert len(world.db.artisan.creates) == 3
    [save] = world.saves
    assert len(save.payload.entries) == 3
    assert report["participantsCreated"] == 3


async def test_an_imported_artisan_is_filed_under_the_design_workshop_and_not_a_field_workshop(
    world,
):
    """``designWorkshopId`` is the column; ``WorkshopArtisan`` is the FIELD workshop's join table.

    Writing that join scopes the artisan to a field workshop that does not exist, and the failure is
    a silent 404 indistinguishable from "no such workshop". The fake client has no
    ``workshopartisan`` delegate at all, so a line that reached for one raises here.
    """
    await run_import(world, [artisan_row()])
    [created] = world.db.artisan.creates
    assert created["designWorkshopId"] == WORKSHOP.id
    assert "workshopId" not in created or created["workshopId"] is None
    assert not hasattr(world.db, "workshopartisan")


async def test_the_roster_goes_through_save_stage_with_the_two_flags_that_matter(world):
    """``replaceCollections=False`` and ``submit=False``, pinned rather than assumed.

    True on the first would sweep every participant row the designer added by hand, and report a
    successful import while doing it. True on the second turns a missing required field into a 422
    on a stage nobody can fix from the screen that caused it.
    """
    await run_import(world, [artisan_row(), artisan_row(name="Sita Bai", aadhaarNumber=AADHAAR_B)])
    [save] = world.saves
    assert save.payload.replaceCollections is False
    assert save.payload.submit is False
    assert save.payload.emptiedEntities == []
    assert save.spec.key == "WORKSHOP_PLAN_PARTICIPANTS_OPENING"
    assert {e.entityKey for e in save.payload.entries} == {"participant"}
    assert [e.data["serialNo"] for e in save.payload.entries] == [1, 2]
    assert [e.ordinal for e in save.payload.entries] == [1, 2]


async def test_a_participant_entry_names_the_artisan_by_reference_and_nothing_else(world):
    """``artisanRef`` plus ``serialNo`` — and no copied fields.

    Copying the artisan's name into the row by hand is what ``hydrate_entries`` does properly, from
    the canonical record, attributing each field to the researcher who recorded it. A hand-copied
    name would be attributed to the officer who uploaded a spreadsheet.
    """
    await run_import(world, [artisan_row()])
    [entry] = world.saves[0].payload.entries
    assert set(entry.data) == {"artisanRef", "serialNo"}
    assert entry.data["artisanRef"] == world.db.artisan.rows[0].id


async def test_a_second_upload_of_the_same_list_adds_no_participant_rows(world):
    """**IDEMPOTENCE, AND IT IS A CORRECTNESS RULE RATHER THAN A SAVING.**

    Re-uploading the same list is the ordinary thing to do after fixing two refused rows, and every
    artisan on it then comes back as `matched` — already recorded, already filed here, nothing
    created. Without the skip the roster would grow a SECOND row for each of them: the printed
    participant table would list fifteen people thirty times, the completeness score would move for
    no reason, and the only way back would be deleting rows by hand on a stage the officer cannot
    open. `replaceCollections=False` is what makes the skip necessary rather than redundant — the
    sweep that would otherwise have replaced the collection is deliberately off.
    """
    world.db.artisan.rows = [existing_artisan(id="artisan-existing", designWorkshopId=WORKSHOP.id)]
    world.db.dwstageentry.rows = [
        SimpleNamespace(
            id="e1",
            designWorkshopId=WORKSHOP.id,
            stageKey="WORKSHOP_PLAN_PARTICIPANTS_OPENING",
            entityKey="participant",
            deletedAt=None,
            ordinal=1,
            data={"artisanRef": "artisan-existing", "serialNo": 1},
        )
    ]
    report = await run_import(world, [artisan_row()])
    assert report["artisansCreated"] == 0
    assert report["artisansLinked"] == 1, "the row was still matched and reported"
    assert report["participantsCreated"] == 0
    assert world.saves == [], "no stage save was issued at all"


async def test_a_second_upload_continues_the_roster_rather_than_colliding_with_it(world):
    """``ordinal`` starts after what the stage already holds, so the printed roster stays in order."""
    world.db.dwstageentry.rows = [
        SimpleNamespace(
            id="e1",
            designWorkshopId=WORKSHOP.id,
            stageKey="WORKSHOP_PLAN_PARTICIPANTS_OPENING",
            entityKey="participant",
            deletedAt=None,
            ordinal=4,
        )
    ]
    await run_import(world, [artisan_row()])
    [entry] = world.saves[0].payload.entries
    assert entry.ordinal == 5
    assert entry.data["serialNo"] == 5


# --------------------------------------------------------------------------------------
# The coordinate
# --------------------------------------------------------------------------------------


async def test_every_imported_artisan_carries_the_venue_fix_and_a_null_subject_pin(world):
    """**THE PROVENANCE / STATED SPLIT, WHICH IS THE MOST IMPORTANT RULE IN THIS FEATURE.**

    The ``Location`` model records the finding that produced it: fifteen artisans documented in four
    different states all carrying one office's coordinates, because the only place the schema
    offered for "where the artisan is" was the field that means "where the device is". The venue fix
    is PROVENANCE — an office keyed the list from the venue — and writing it into the subject pin
    would be a claim about where the artisan lives. Both columns are floats and nothing refuses it,
    which is why this is a test.
    """
    await run_import(world, [artisan_row()])
    [location] = world.db.location.creates
    assert (location["latitude"], location["longitude"]) == (VENUE_FIX["lat"], VENUE_FIX["lon"])
    assert location["accuracy"] == VENUE_FIX["accuracy"]
    assert location.get("subjectLatitude") is None
    assert location.get("subjectLongitude") is None
    assert location["placeName"] == "Bagru"
    assert location.get("address") is None, "Location.address is the device's reverse geocode"


async def test_an_imported_location_says_it_came_from_an_import(world):
    """The only way a later reader can tell fifteen artisans sharing one fix from fifteen
    researchers who happened to stand in one place."""
    report = await run_import(world, [artisan_row()])
    [location] = world.db.location.creates
    meta = stored(location["extraMetadata"])
    assert meta["source"] == "ARTISAN_XLSX_IMPORT"
    assert meta["designWorkshopId"] == WORKSHOP.id
    assert meta["importId"] == report["importId"]


async def test_the_import_metadata_keys_cannot_be_read_as_a_stated_address(world):
    """``schemas/common._stated_district`` reads ``extraMetadata["district"]`` as a FALLBACK.

    A provenance key by that name — or by ``state``, ``village`` or ``pincode`` — would silently
    become a stated address nobody typed, and ``require_location`` would pass on it.
    """
    await run_import(world, [artisan_row()])
    meta = stored(world.db.location.creates[0]["extraMetadata"])
    assert not ({"district", "state", "village", "pincode"} & set(meta))


async def test_the_row_supplies_the_stated_address_and_the_workshop_supplies_the_coordinate(world):
    await run_import(world, [artisan_row(village="Bagru Khurd", pincode="303007")])
    [location] = world.db.location.creates
    assert location["state"] == "Rajasthan"
    assert location["district"] == "Jaipur"
    assert location["village"] == "Bagru Khurd"
    assert location["pincode"] == "303007"


# --------------------------------------------------------------------------------------
# Never duplicate, never overwrite
# --------------------------------------------------------------------------------------


def existing_artisan(**overrides):
    base = {
        "id": "artisan-existing",
        "name": "Ramesh Devi",
        "place": "Bagru",
        "phone": "9000000000",
        "email": None,
        "aadhaarNumber": AADHAAR_A,
        "pehchanCardNumber": None,
        "designWorkshopId": "cmworkshop00000000000099",
        "designWorkshop": SimpleNamespace(title="An earlier workshop"),
        "craft": SimpleNamespace(name="Block Printing"),
    }
    base.update(overrides)
    return SimpleNamespace(**base)


async def test_an_existing_artisan_is_linked_and_not_duplicated(world):
    world.db.artisan.rows = [existing_artisan()]
    report = await run_import(world, [artisan_row()])
    assert report["artisansCreated"] == 0
    assert report["artisansLinked"] == 1
    assert world.db.artisan.creates == [], "no artisan write at all"
    assert world.db.artisan.updates == [], "and no update either"
    [entry] = world.saves[0].payload.entries
    assert entry.data["artisanRef"] == "artisan-existing"


async def test_a_linked_artisan_is_not_refiled_into_this_workshop(world):
    """Setting ``designWorkshopId`` on somebody else's artisan moves that record out of THEIR
    workshop's scoped lists and totals. The participant row is the link this workshop needs."""
    existing = existing_artisan()
    world.db.artisan.rows = [existing]
    await run_import(world, [artisan_row()])
    assert existing.designWorkshopId == "cmworkshop00000000000099"
    assert world.db.artisan.updates == []


async def test_an_artisan_already_on_this_workshop_is_reported_as_matched_rather_than_linked(world):
    world.db.artisan.rows = [existing_artisan(designWorkshopId=WORKSHOP.id, designWorkshop=None)]
    report = await run_import(world, [artisan_row()])
    assert report["artisansLinked"] == 1
    assert any("already on this workshop" in p["reason"] for p in report["problems"])


async def test_a_disagreeing_field_is_reported_and_the_stored_value_stands(world):
    """**NEVER SILENTLY OVERWRITE.** ``merge_artisans.py`` is what correcting a bad merge costs."""
    existing = existing_artisan(place="Sanganer", phone="9111111111")
    world.db.artisan.rows = [existing]
    report = await run_import(world, [artisan_row(place="Bagru", phone="9222222222")])
    reasons = " ".join(p["reason"] for p in report["problems"])
    assert "Place differs" in reasons
    assert "Phone differs" in reasons
    assert "Sanganer" in reasons and "Bagru" in reasons
    assert existing.place == "Sanganer", "the stored record was not changed"
    assert world.db.artisan.updates == []


async def test_two_rows_with_one_aadhaar_import_the_first_and_refuse_the_second_by_row_number(
    world,
):
    """An error and NOT the warning the questionnaire parser gives for a duplicate question id.

    ``Artisan.aadhaarNumber`` is ``@unique``, so importing the second row would take the index
    mid-batch as a bare 500 after half the file was written.
    """
    report = await run_import(
        world, [artisan_row(), artisan_row(name="Ramesh Devi again", aadhaarNumber=AADHAAR_A)]
    )
    assert report["artisansCreated"] == 1
    assert report["rowsRefused"] == 1
    refusal = next(p for p in report["problems"] if p["severity"] == "error")
    assert refusal["row"] == 3
    assert "row 2" in refusal["reason"], "it names the EARLIER row"
    assert "cannot be recorded twice" in refusal["reason"]


async def test_a_row_that_loses_the_unique_race_is_linked_rather_than_failing_the_batch(world):
    """A fifteen-row upload must not be lost to one contested row.

    The row becomes ``linked``, which is the same answer the pre-flight would have given a second
    earlier — not a 409, because this is a BATCH.
    """
    # The winner is written by "somebody else" AFTER the pre-flight has run and BEFORE the create
    # lands — which is exactly the window the unique index exists to close. The fake reproduces it
    # by raising the driver's message once and having the row already present for the re-read.
    world.db.artisan.create_error = Exception(
        "Unique constraint failed on the fields: (`aadhaarNumber`)"
    )
    world.db.artisan.rows = [existing_artisan(id="artisan-winner", aadhaarNumber=AADHAAR_B)]
    parsed = parse_artisan_workbook(
        sheet([artisan_row(aadhaarNumber=AADHAAR_B)]), defaults=DEFAULTS
    )
    # Emptied so the pre-flight does NOT see the winner, which is what makes this the RACE path
    # rather than the ordinary "already recorded" one.
    import app.services.artisan_import as importer

    async def no_clashes(_rows):
        return {}

    original, importer._identity_clashes = importer._identity_clashes, no_clashes
    try:
        report = await import_artisans(
            parsed,
            workshop=WORKSHOP,
            actor=OFFICER,
            venue_fix=VENUE_FIX,
            source_filename="race.xlsx",
        )
    finally:
        importer._identity_clashes = original

    assert report["rowsRefused"] == 0, "a fifteen-row upload must not be lost to one contested row"
    assert report["artisansCreated"] == 0
    assert world.db.artisan.creates == [], "the create was refused by the index, not retried"
    [entry] = world.saves[0].payload.entries
    assert entry.data["artisanRef"] == "artisan-winner", "the row links to whoever won the race"


async def test_a_race_on_a_different_column_is_raised_rather_than_silently_linked(world):
    """Only a unique violation on an IDENTITY column means "somebody else recorded this person"."""
    world.db.artisan.create_error = Exception("some other database failure")
    with pytest.raises(Exception, match="some other database failure"):
        await run_import(world, [artisan_row()])


# --------------------------------------------------------------------------------------
# The craft
# --------------------------------------------------------------------------------------


async def test_a_craft_name_matching_nothing_is_refused_and_no_craft_is_created(world):
    """``allow_create=False``. A typo in one cell must not mint a row in the controlled vocabulary.

    The actor here CLEARS ``can_manage_crafts`` — every officer tier outranks PROFESSOR — so the
    refusal is a decision about SPREADSHEETS and not about rank, which is exactly why the parameter
    exists.
    """
    from app.core.deps import can_manage_crafts

    assert can_manage_crafts(OFFICER), "the control: this actor could create a craft on the form"
    report = await run_import(world, [artisan_row(craftName="Blok printing")])
    assert report["artisansCreated"] == 0
    assert report["rowsRefused"] == 1
    assert world.db.craft.creates == [], "the controlled vocabulary is unchanged"
    refusal = next(p for p in report["problems"] if p["severity"] == "error")
    # TITLE-CASED, because ``clean_data`` normalises ``craftName`` before the lookup sees it — the
    # same spelling ``POST /crafts`` would have stored. The refusal quotes what was LOOKED FOR.
    assert "Blok Printing" in refusal["reason"]
    assert "controlled list" in refusal["reason"]


async def test_a_known_craft_is_resolved_to_its_id(world):
    await run_import(world, [artisan_row()])
    [created] = world.db.artisan.creates
    assert created["craftId"] == "craft-1"
    assert "craftName" not in created


# --------------------------------------------------------------------------------------
# The review status — measured, and the plan's assumption was wrong
# --------------------------------------------------------------------------------------


async def test_an_imported_artisan_is_created_pending_even_though_the_officer_outranks_professor(
    world,
):
    """**THE MEASURED ANSWER, PINNED IN BOTH DIRECTIONS.**

    ``apply_status_policy_create`` reads "Professor and above default to APPROVED", and every
    officer tier clears that floor — so the obvious conclusion is that fifteen never-reviewed
    artisans enter the repository approved with the review queue showing nothing. **They do not**,
    because that function uses ``setdefault`` and ``ArtisanCreate.status`` is declared
    ``str = "PENDING"``: the key is already in the dict by the time the policy runs.

    The second half of this test is what makes the first half meaningful. Without it, deleting the
    schema default would leave this test passing for a completely different reason — every artisan
    silently approved — which is the outcome the whole assertion exists to notice.
    """
    from app.services.records import apply_status_policy_create

    await run_import(world, [artisan_row()])
    [created] = world.db.artisan.creates
    assert created["status"] == "PENDING"

    # The control: the senior branch IS reachable, by a caller that omits the key.
    assert apply_status_policy_create(OFFICER, {})["status"] == "APPROVED"


# --------------------------------------------------------------------------------------
# The ledger, and the identity numbers in it
# --------------------------------------------------------------------------------------


async def test_the_ledger_row_is_written_before_any_artisan(world):
    """So a deployment whose migration has not run 500s having changed nothing.

    It is also what every ``Location``'s provenance points at, which is the other reason it cannot
    be written at the end.
    """
    await run_import(world, [artisan_row()])
    assert world.db.dwartisanimport.creates, "the ledger row exists"
    assert world.db.dwartisanimport.updates, "and its counts were filled in afterwards"
    created = world.db.dwartisanimport.creates[0]
    assert created["sourceFilename"] == "bagru-list.xlsx"
    assert created["designWorkshopId"] == WORKSHOP.id
    assert created["uploadedById"] == OFFICER.id


async def test_the_workbook_itself_is_never_stored_only_its_filename(world):
    """One regulated COLUMN must not become a regulated FILE with its own retention question."""
    await run_import(world, [artisan_row()])
    created = world.db.dwartisanimport.creates[0]
    assert set(created) == {
        "designWorkshopId",
        "uploadedById",
        "sourceFilename",
        "sheetName",
        "problems",
    }


async def test_the_import_is_recorded_in_the_ledger_with_masked_problems(world):
    """**No twelve-digit run reaches the ledger.** Asserted over the whole stored blob.

    A diagnostic nobody reads as a data path is still a data path — and unlike the HTTP response,
    this one is kept for the life of the workshop.
    """
    import re

    report = await run_import(
        world,
        [
            artisan_row(),
            artisan_row(name="Bad", aadhaarNumber="223456789019"),  # valid shape, bad checksum
            artisan_row(name="Twin", aadhaarNumber=AADHAAR_A),
        ],
    )
    (_where, data) = world.db.dwartisanimport.updates[-1]
    blob = str(stored(data["problems"]))
    assert re.search(r"\d{12}", blob) is None, blob
    assert "XXXX XXXX" in blob, "the masked form IS carried, so the row is identifiable"
    assert data["rowsRead"] == 3
    assert data["artisansCreated"] == 1
    assert data["rowsRefused"] == 2
    assert blob == str(report["problems"]), "the ledger and the response carry the same list"


@pytest.mark.parametrize("flag", ["", "No", "Yes"])
async def test_a_masked_pehchan_row_is_refused_rather_than_created_without_its_card_number(
    world, flag
):
    """**THE DEFECT THE LEDGER IS THE PERMANENT RECORD OF.**

    ``_read_pehchan`` filed a ``severity="error"`` for a card number pasted off a screen and then
    let the row through, and ``artisan_import`` counts ``refused`` off the problem list:
    ``refused = len({p.row for p in problems if p.severity == "error"})``. So ONE row was counted in
    both ``artisansCreated`` and ``rowsRefused``, the balance check logged "does not balance" — a
    line the module marks ``# pragma: no cover - a bug, not a state`` — and the officer was told a
    row had been refused that was in fact in the repository, with its regulated card number silently
    dropped.

    Only the ``Yes`` branch was ever caught (``reconcile_pehchan`` demands a number beside it), and
    it balanced by accident because its second error landed on the same row and ``refused_rows`` is
    a set. The two states an office actually produces — the flag column left alone, or answered No —
    were not caught at all, and re-uploading the corrected sheet does not repair them: the row
    matches on Aadhaar and takes the branch whose own message is "the existing record was not
    changed and was not re-filed".
    """
    report = await run_import(
        world,
        [artisan_row(pehchanCardAvailable=flag, pehchanCardNumber="XXXX XXXX 3456")],
    )
    assert report["rowsRead"] == 1
    assert report["rowsRefused"] == 1
    assert report["artisansCreated"] == 0, (
        "the artisan was created after the report said the row was refused, and without the card "
        "number the sheet was trying to record"
    )
    assert world.db.artisan.creates == []
    refusal = next(p for p in report["problems"] if p["severity"] == "error")
    assert "masked" in refusal["reason"]


async def test_every_report_balances_rows_read_against_what_became_of_each_row(world):
    """``created + linked + refused == rowsRead``, asserted rather than logged.

    ``import_artisans`` checks this itself and answers a ``logger.warning`` when it fails — nothing
    reaches a screen, and the 201 goes out with numbers that do not add up. The mixed sheet below is
    the one that used to fail it: a good row, a row the parser refuses outright, and a row the
    parser used to call refused and keep.
    """
    report = await run_import(
        world,
        [
            artisan_row(),
            artisan_row(name="No name at all", aadhaarNumber=AADHAAR_B, state="Tamil Nady"),
            artisan_row(
                name="Masked card",
                aadhaarNumber=AADHAAR_C,
                pehchanCardNumber="XXXX XXXX 3456",
            ),
        ],
    )
    assert (
        report["artisansCreated"] + report["artisansLinked"] + report["rowsRefused"]
        == report["rowsRead"]
    ), report


async def test_the_counts_in_the_ledger_and_the_response_agree(world):
    report = await run_import(
        world, [artisan_row(), artisan_row(name="Sita Bai", aadhaarNumber=AADHAAR_B)]
    )
    (_where, data) = world.db.dwartisanimport.updates[-1]
    for key in (
        "rowsRead",
        "artisansCreated",
        "artisansLinked",
        "rowsRefused",
        "participantsCreated",
    ):
        assert data[key] == report[key], key


# --------------------------------------------------------------------------------------
# The invariant the roster write rests on
# --------------------------------------------------------------------------------------


def test_the_stage_chunk_stays_under_the_stage_row_bound():
    """Compared LIVE, never against two literals.

    A ``StageSaveIn`` built IN PROCESS that exceeds ``MAX_STAGE_ROWS`` raises a
    ``pydantic.ValidationError`` rather than a 422 — there is no request-body wrapper around it — and
    that becomes a bare 500 AFTER every artisan in the file is already committed. A test written
    against the numbers would go green the day somebody lowered the bound.
    """
    from app.schemas.design_workshops import MAX_STAGE_ROWS
    from app.services.artisan_import import STAGE_CHUNK
    from app.services.artisan_xlsx import MAX_ARTISANS

    assert STAGE_CHUNK <= MAX_STAGE_ROWS
    assert MAX_ARTISANS <= MAX_STAGE_ROWS


def test_the_importer_writes_no_stage_entry_of_its_own():
    """**THE THIRD-WRITER RULE, ASSERTED AT THE SOURCE.**

    ``tests/test_design_workshop_search_text.py`` already sweeps ``backend/app`` for writers of
    ``DwStageEntry`` and asserts the register is EXACTLY two, both of them in
    ``services/design_workshops``. This is the same claim pointed at this one file, so that a
    failure names the module that caused it rather than a register somebody might be tempted to
    widen.

    A row written without ``searchText`` is invisible to ``GET /search`` for the workshop's whole
    life: NULL there means "not computed yet" and can never match.
    """
    import ast
    from pathlib import Path

    source = Path(__file__).resolve().parents[1] / "app" / "services" / "artisan_import.py"
    tree = ast.parse(source.read_text(encoding="utf-8"))
    writes = [
        node
        for node in ast.walk(tree)
        if isinstance(node, ast.Call)
        and isinstance(node.func, ast.Attribute)
        and isinstance(node.func.value, ast.Attribute)
        and node.func.value.attr == "dwstageentry"
        and node.func.attr
        in {"create", "create_many", "update", "update_many", "upsert", "delete_many"}
    ]
    assert writes == [], (
        "artisan_import has grown a direct DwStageEntry write. Read this module's header: the "
        "roster goes through design_workshops.save_stage, which is what writes searchText, "
        "hydrates the ~30 reference fields, holds the version guard and attributes provenance to "
        "the artisan's own recorder."
    )
