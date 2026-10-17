"""C11 — a designer imports their own workshop's artisan list, through the officers' one pipeline.

Owner decision (2026-10-17): designers may import their own workshop's roster, under the same
delete-after-reading and audit rules. The import itself is covered by ``test_artisan_import``; this
file holds the DOOR: who it admits, that it is the same pipeline, and that it records the ledger.

Arranged like ``tests/test_sanction_orders.py``: database calls only in private ``asyncio.run``
loops while no ``TestClient`` is alive.
"""

from __future__ import annotations

import asyncio
import random
import uuid
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from typing import Any

import pytest
from conftest import needs_db
from openpyxl import load_workbook

from app.api.routes import design_workshop_roster
from app.core.db import db
from app.core.security import create_access_token, hash_password
from app.services.artisan_identity import verhoeff_ok
from app.services.artisan_xlsx import _COLUMNS, build_artisan_pro_forma
from prisma import Json

ROUTES = Path(design_workshop_roster.__file__).resolve()


def test_the_designer_door_is_the_stage_write_door_and_the_officers_pipeline() -> None:
    source = ROUTES.read_text(encoding="utf-8")
    assert source.count("for_edit=True") == 2
    assert source.count("_require_designer(current_user)") == 2
    assert "run_artisan_import(" in source and "artisan_import_history(" in source
    # One pipeline: the import service is never called from the designer's module directly.
    assert "import_artisans" not in source


def test_the_router_is_mounted() -> None:
    from app.main import app

    paths = {getattr(route, "path", "") for route in app.routes}
    assert "/api/design-workshops/{workshop_id}/roster/upload" in paths
    assert "/api/design-workshops/{workshop_id}/roster/pro-forma" in paths
    assert "/api/design-workshops/{workshop_id}/roster/imports" in paths


def _aadhaar() -> str:
    while True:
        stem = str(random.randint(2, 9)) + "".join(str(random.randint(0, 9)) for _ in range(10))
        for check in "0123456789":
            if verhoeff_ok(stem + check):
                return stem + check


def _workbook(workshop_id: str, craft: str) -> bytes:
    order = [key for key, _label in _COLUMNS]
    wb = load_workbook(BytesIO(build_artisan_pro_forma()))
    ws = wb["Artisans"]
    values = {
        "name": "Sita Devi",
        "craftName": craft,
        "place": "Ajrakhpur",
        "aadhaarNumber": _aadhaar(),
        "dos": "Ask before photographing",
        "donts": "Do not hurry the dye",
        "state": "Gujarat",
        "district": "Kachchh",
    }
    for column, key in enumerate(order, start=1):
        if values.get(key):
            ws.cell(row=2, column=column, value=values[key])
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


@pytest.fixture(scope="module")
def world() -> Any:
    from fastapi.testclient import TestClient

    from app.main import app

    stamp = uuid.uuid4().hex[:8]
    facts: dict[str, Any] = {"stamp": stamp}
    craft_name = f"Roster craft {stamp}"

    async def seed() -> None:
        await db.connect()
        try:
            for slug, role in (
                ("owner", "ADMIN"),
                ("stranger", "DESIGNER"),
                ("reader", "RESEARCHER"),
            ):
                email = f"c11-{slug}-{stamp}@example.org"
                facts[slug] = await db.user.create(
                    data={
                        "email": email,
                        "name": f"C11 {slug} {stamp}",
                        "role": role,
                        "passwordHash": hash_password("c11-password-123"),
                    }
                )
                await db.accessroster.create(
                    data={
                        "email": email,
                        "status": "ACTIVE",
                        "admitRole": role,
                        "joinedAt": datetime.now(UTC),
                        "notes": "Seeded by tests/test_design_workshop_roster_import.py.",
                    }
                )
            await db.craft.create(data={"name": craft_name})
            facts["bare"] = await db.designworkshop.create(
                data={
                    "title": f"No venue {stamp}",
                    "templateId": "DCH_STANDARD",
                    "createdById": facts["owner"].id,
                }
            )
            ready = await db.designworkshop.create(
                data={
                    "title": f"With venue {stamp}",
                    "templateId": "DCH_STANDARD",
                    "createdById": facts["owner"].id,
                    "craftName": craft_name,
                    "state": "Gujarat",
                    "district": "Kachchh",
                }
            )
            await db.dwstageentry.create(
                data={
                    "designWorkshopId": ready.id,
                    "stageKey": "WORKSHOP_SETUP",
                    "entityKey": "workshopSetup",
                    "data": Json({"venueLocation": {"lat": 23.21, "lon": 69.67, "accuracy": 10.0}}),
                }
            )
            facts["ready"] = ready
        finally:
            await db.disconnect()

    asyncio.run(seed())

    def headers(slug: str) -> dict[str, str]:
        return {"Authorization": f"Bearer {create_access_token(subject=facts[slug].id)}"}

    def upload(workshop_id: str, slug: str):
        return client.post(
            f"/api/design-workshops/{workshop_id}/roster/upload",
            files={
                "file": (
                    "roster.xlsx",
                    _workbook(workshop_id, craft_name),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            headers=headers(slug),
        )

    with TestClient(app) as client:
        facts["no_venue"] = upload(facts["bare"].id, "owner")
        facts["stranger_upload"] = upload(facts["ready"].id, "stranger")
        facts["reader_upload"] = upload(facts["ready"].id, "reader")
        facts["pro_forma"] = client.get(
            f"/api/design-workshops/{facts['ready'].id}/roster/pro-forma", headers=headers("owner")
        )
        facts["imported"] = upload(facts["ready"].id, "owner")
        facts["history"] = client.get(
            f"/api/design-workshops/{facts['ready'].id}/roster/imports", headers=headers("owner")
        )

    async def observe() -> None:
        await db.connect()
        try:
            facts["ledger"] = await db.dwartisanimport.find_many(
                where={"designWorkshopId": facts["ready"].id}
            )
            facts["artisans"] = await db.artisan.find_many(
                where={"designWorkshopId": facts["ready"].id}
            )
        finally:
            await db.disconnect()

    asyncio.run(observe())
    yield facts


@needs_db
def test_a_workshop_with_no_venue_is_refused_by_the_shared_pipeline(world: Any) -> None:
    assert world["no_venue"].status_code == 422
    assert "venue location" in world["no_venue"].json()["detail"]


@needs_db
def test_somebody_not_on_the_workshop_cannot_import_into_it(world: Any) -> None:
    assert world["stranger_upload"].status_code == 404
    assert world["reader_upload"].status_code == 403


@needs_db
def test_the_pro_forma_names_the_workshop(world: Any) -> None:
    assert world["pro_forma"].status_code == 200
    wb = load_workbook(BytesIO(world["pro_forma"].content))
    details = "\n".join(
        str(cell.value) for row in wb.worksheets[-1].iter_rows() for cell in row if cell.value
    ) + "\n".join(
        str(cell.value)
        for sheet in wb.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.value
    )
    assert world["ready"].id in details


@needs_db
def test_the_owner_imports_and_the_ledger_records_it(world: Any) -> None:
    assert world["imported"].status_code == 201, world["imported"].text
    report = world["imported"].json()
    assert report["artisansCreated"] == 1
    assert report["participantsCreated"] == 1
    assert len(world["ledger"]) == 1
    assert world["ledger"][0].uploadedById == world["owner"].id
    assert world["ledger"][0].sourceFilename == "roster.xlsx"
    assert [artisan.name for artisan in world["artisans"]] == ["Sita Devi"]
    history = world["history"].json()["imports"]
    assert history[0]["uploadedBy"] == world["owner"].name
