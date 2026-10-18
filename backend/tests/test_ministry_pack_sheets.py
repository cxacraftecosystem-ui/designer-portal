"""THE MINISTRY PACK'S SHEETS, WITHOUT A DATABASE.

``services/ministry_pack_sheets`` is pure: rows in, sheet dicts out. These tests hold what a file that
leaves the product must and must not say — Summary first, one sheet per register, no contact details,
no twelve-digit runs, no links, no soft-deleted workshops, UNSTATED counted, a truncation note when
the cap was hit — and that the workbook it renders opens cleanly.
"""

from __future__ import annotations

import re
import xml.etree.ElementTree as ET
import zipfile
from datetime import UTC, datetime
from decimal import Decimal
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from app.services import ministry_pack_sheets as sheets
from app.services.mailer import IST
from app.services.xlsx_report import build_report_workbook

START, END = sheets.period_bounds("2026-09", IST)


def _workshop(index: int, **extra: Any) -> dict[str, Any]:
    row = {
        "id": f"w{index}",
        "title": f"Workshop {index}",
        "status": "IN_PROGRESS",
        "workshopCode": f"DPW-{index:03d}",
        "workshopKind": "SKILL_UPGRADATION",
        "state": "Rajasthan" if index % 2 else "Odisha",
        "participants": 10,
        "createdAt": datetime(2026, 9, 10, tzinfo=UTC),
        "updatedAt": datetime(2026, 9, 12, tzinfo=UTC),
        "deletedAt": None,
    }
    row.update(extra)
    return row


def _data(workshops: list[dict[str, Any]], **extra: Any) -> dict[str, Any]:
    data = {
        "period": "2026-09",
        "periodLabel": "September 2026",
        "scheme": "SKILL_UPGRADATION",
        "schemeLabel": "Skill Upgradation",
        "snapshotAt": datetime(2026, 10, 1, 1, 0, tzinfo=UTC),
        "monthToDate": False,
        "revision": 1,
        "truncated": False,
        "cap": 2000,
        "periodStart": START,
        "periodEnd": END,
        "workshops": workshops,
        "progress": {},
        "feedback": {},
        "exports": {},
        "annualPlan": [],
        "planYearLabel": "2026-27",
    }
    data.update(extra)
    return data


def test_periods_labels_and_financial_years() -> None:
    assert sheets.period_label("2026-09") == "September 2026"
    assert sheets.parse_period("2026-13") is None
    assert sheets.parse_period("26-09") is None
    assert sheets.plan_year_for("2026-03") == 2025
    assert sheets.plan_year_for("2026-04") == 2026
    start, end = sheets.period_bounds("2025-12", IST)
    assert start == datetime(2025, 11, 30, 18, 30, tzinfo=UTC)
    assert end == datetime(2025, 12, 31, 18, 30, tzinfo=UTC)


def test_scheme_keys_are_the_six_kinds_then_unstated() -> None:
    keys = sheets.scheme_keys()
    assert len(keys) == 7 and keys[-1] == sheets.UNSTATED
    assert sheets.scheme_label(sheets.UNSTATED) == "Scheme not yet stated"
    assert sheets.scheme_label("SKILL_UPGRADATION") == "Skill Upgradation"


def test_a_sheet_per_register_with_summary_first() -> None:
    built, counts, _ = sheets.build_sheets(_data([_workshop(1), _workshop(2)]))
    assert [s["name"] for s in built] == [
        "Summary",
        "Workshops",
        "Sanction orders",
        "Annual plan",
        "Inspections",
        "Report hand-ons",
    ]
    assert counts["workshops"] == 2 and counts["participants"] == 20


def test_counts_come_from_real_timestamps_inside_the_period() -> None:
    rows = [
        _workshop(
            1, status="DRAFT", createdAt=datetime(2026, 8, 31, 19, 0, tzinfo=UTC)
        ),  # 1 Sep IST
        _workshop(2, status="APPROVED", approvedAt=datetime(2026, 9, 30, 18, 29, tzinfo=UTC)),
        _workshop(
            3, status="APPROVED", approvedAt=datetime(2026, 9, 30, 18, 31, tzinfo=UTC)
        ),  # 1 Oct IST
        _workshop(
            4,
            status="SUBMITTED",
            lastHandedInAt=datetime(2026, 9, 5, tzinfo=UTC),
            sanctionAmount=Decimal("1234500.00"),
        ),
    ]
    counts = sheets.compute_counts(rows, START, END, truncated=False)
    assert counts["registered"] == 1 and counts["completed"] == 3 and counts["ongoing"] == 0
    assert counts["approvedInPeriod"] == 1
    assert counts["handedInInPeriod"] == 1
    assert counts["createdInPeriod"] == 4
    assert counts["sanctionedAmount"] == "1234500.00"
    assert sheets.indian_number(counts["sanctionedAmount"]) == "12,34,500"


def test_the_scrub_masks_identity_and_drops_contact_and_links() -> None:
    sheet = {
        "name": "Stage",
        "color": "#000000",
        "columns": ["Name", "Phone", "Email", "Aadhaar", "Note", "Photo"],
        "keys": ["p.name", "p.phone", "p.email", "p.aadhaarNumber", "p.note", "p.photo"],
        "rows": [
            [
                "Asha",
                "9876543210",
                "a@example.org",
                "123456789012",
                "card 2345 6789 0123 seen",
                "https://bucket/x?X-Amz-Signature=1",
            ],
            ["Ravi", None, None, "XXXX XXXX 4321", "plain", "http://example.org"],
        ],
    }
    clean = sheets.scrub_sheet(sheet)
    assert clean["columns"] == ["Name", "Aadhaar", "Note", "Photo"]
    assert clean["rows"][0][1] == "XXXX XXXX 9012"
    assert clean["rows"][1][1] == "XXXX XXXX 4321"
    assert clean["rows"][0][2] == "card XXXX XXXX 0123 seen"
    flat = [str(c) for c in sheets.all_cells([clean]) if c is not None]
    assert not any(re.search(r"(?<!\d)\d{12}(?!\d)", c.replace(" ", "")) for c in flat)
    assert not any("http" in c or "X-Amz" in c for c in flat)
    for row in clean["rows"]:
        assert re.fullmatch(r"X{4} X{4} \d{4}", row[1])


def test_the_scrub_matches_labels_when_a_sheet_has_no_keys() -> None:
    clean = sheets.scrub_sheet(
        {
            "name": "x",
            "columns": ["Designer", "Designer e-mail", "Mobile", "Pehchan card"],
            "rows": [["A", "a@b.c", "98", "PC12345678"]],
        }
    )
    assert clean["columns"] == ["Designer", "Pehchan card"]
    assert clean["rows"][0] == ["A", "XXXX XXXX 5678"]


def test_no_cell_in_a_pack_carries_a_link_a_signature_or_twelve_digits() -> None:
    rows = [
        _workshop(1, venue="Hall, call 1234 5678 9012", notes="https://x"),
        _workshop(2, title="https://evil.example/X-Amz-Credential", status="SUBMITTED"),
    ]
    built, _, _ = sheets.build_sheets(_data(rows))
    flat = [str(c) for c in sheets.all_cells(built) if c is not None]
    assert not any("X-Amz" in c or c.lower().startswith("http") for c in flat)
    assert not any(
        re.search(r"(?<![0-9A-Za-z])\d{4}[ -]?\d{4}[ -]?\d{4}(?![0-9A-Za-z])", c) for c in flat
    )


def test_a_mobile_number_or_address_typed_into_free_text_is_removed() -> None:
    rows = [
        _workshop(1, venue="Hall, call 98765 43210", title="Ask +91-9876543210 or a.b@example.org"),
        _workshop(2, venue="Call 09876543210", title="Block 12, batch 2026"),
    ]
    built, _, _ = sheets.build_sheets(_data(rows))
    flat = [str(c) for c in sheets.all_cells(built) if c is not None]
    assert not any(re.search(r"[6-9]\d{4}[ -]?\d{5}", c) for c in flat)
    assert not any("@example.org" in c for c in flat)
    assert "Hall, call [phone removed]" in flat
    assert "Block 12, batch 2026" in flat
    assert sheets.scrub_sheet({"columns": ["Note"], "rows": [["ring 0 98765 43210"]]})["rows"] == [
        ["ring 0 [phone removed]"]
    ]


def test_no_phone_or_email_column_survives() -> None:
    built, _, _ = sheets.build_sheets(
        _data([_workshop(1, sanctionOrderNo="SO-1", sanctionAmount="10")])
    )
    headers = [str(h).lower() for s in built for h in s["columns"]]
    assert not any("phone" in h or "mail" in h for h in headers)


def test_soft_deleted_workshops_are_excluded() -> None:
    built, counts, _ = sheets.build_sheets(
        _data([_workshop(1), _workshop(2, deletedAt=datetime(2026, 9, 20, tzinfo=UTC))])
    )
    assert counts["workshops"] == 1
    register = next(s for s in built if s["name"] == "Workshops")
    assert [r[0] for r in register["rows"]] == ["DPW-001"]


def test_draft_and_in_progress_workshops_are_in_the_register_without_stage_answers() -> None:
    built, _, _ = sheets.build_sheets(
        _data([_workshop(1, status="DRAFT"), _workshop(2, status="IN_PROGRESS")])
    )
    register = next(s for s in built if s["name"] == "Workshops")
    assert {r[11] for r in register["rows"]} == {"DRAFT", "IN_PROGRESS"}
    # The pack carries registers only: no sheet is built from stage answers.
    assert all(
        s["name"]
        in {
            "Summary",
            "Workshops",
            "Sanction orders",
            "Annual plan",
            "Inspections",
            "Report hand-ons",
        }
        for s in built
    )


def test_unstated_is_labelled_and_reads_from_a_null_kind() -> None:
    data = _data(
        [_workshop(1, workshopKind=None)],
        scheme=sheets.UNSTATED,
        schemeLabel=sheets.scheme_label(sheets.UNSTATED),
    )
    built, counts, _ = sheets.build_sheets(data)
    summary = built[0]
    assert ["Scheme", "Scheme not yet stated", None, None] in summary["rows"]
    assert counts["workshops"] == 1


def test_truncation_and_month_to_date_are_written_into_the_summary() -> None:
    built, counts, _ = sheets.build_sheets(
        _data([_workshop(1)], truncated=True, cap=1, monthToDate=True)
    )
    labels = [row[0] for row in built[0]["rows"]]
    assert counts["truncated"] is True
    assert "Truncated" in labels and "Month to date" in labels
    assert "Data protection" in labels


def test_the_scope_sentence_names_the_scheme() -> None:
    built, _, _ = sheets.build_sheets(_data([_workshop(1)]))
    assert ["Scope", "Every workshop under Skill Upgradation on the platform", None, None] in built[
        0
    ]["rows"]


def test_by_state_puts_the_biggest_first_and_unnamed_last() -> None:
    table = sheets.by_state([_workshop(1), _workshop(3), _workshop(2), _workshop(5, state=None)])
    assert [e["state"] for e in table] == ["Rajasthan", "Odisha", "State not given"]


def test_the_workbook_opens_cleanly_with_summary_first() -> None:
    rows = [
        _workshop(
            i,
            title=f"=SUM(A{i}) weaver's tools",
            sanctionOrderNo=f"SO-{i}",
            sanctionAmount="1500.5",
        )
        for i in range(1, 6)
    ]
    plan = [{"workshopNo": "1", "plannedTitle": "Plan", "designWorkshopId": "w1"}]
    built, _, _ = sheets.build_sheets(_data(rows, annualPlan=plan))
    body = build_report_workbook(built, "Ministry pack", overview=False)
    archive = zipfile.ZipFile(BytesIO(body))
    for name in archive.namelist():
        if name.endswith((".xml", ".rels")):
            ET.fromstring(archive.read(name))
    wb = load_workbook(BytesIO(body))
    assert wb.sheetnames[0] == "Summary"
    assert "Overview" not in wb.sheetnames
    register = wb["Workshops"]
    assert register.cell(row=2, column=2).data_type == "s"  # the "=" title stayed text
