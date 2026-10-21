"""The Ministry dashboard's stage-data export (owner decision C15, 2026-10-17).

One spreadsheet, a sheet per record type (the entity registry ``design_workshop_data.tables()``),
across the workshops the requester's dashboard filters and scope select — with every identity
number masked — built in the BACKGROUND by the queue worker, because a few hundred workshops of
stage rows is minutes of openpyxl on a box the API shares.

WHO. The Ministry dashboard's readers (``deps.MINISTRY_DASHBOARD_ROLES``): Assistant and Regional
Directors over the workshops they were posted to, the Ministry Admin and the master admin over the
estate. The scope is re-read from the requester's account when the job RUNS, so a posting removed
between the request and the run narrows the file rather than leaking it.

WHAT IS LEFT OUT, AND SAID ON THE README SHEET. Media are cells of ids (never bytes or URLs, the
registry's own export rule). A designer's own custom questions are withheld: their columns are
defined per workshop and nothing here can tell which of them holds an identity number. Every field
the registry marks as an Aadhaar or Pehchan number, or whose key names one, is masked to its last
four characters — including values a designer typed in full.

WHERE THE FILE LIVES. In the ``MinistryDataExport`` row, for :data:`EXPORT_TTL_DAYS`, then purged by
the worker. Only the account that asked may download it.
"""

from __future__ import annotations

import asyncio
import json
import logging
import re
from datetime import UTC, datetime, timedelta
from io import BytesIO
from typing import Any

from app.core.db import db
from app.services import design_workshop_data as dw, naming
from app.services.records import mask_identity_number
from app.services.stage_schema import TextFormat, stages

log = logging.getLogger(__name__)

#: How long a finished file is kept before the worker purges it.
EXPORT_TTL_DAYS = 7
#: How many workshops one export reads. Past it the README says the file stopped.
MAX_EXPORT_WORKSHOPS = 1000
#: How many stage rows one export writes, across every sheet.
MAX_EXPORT_ROWS = 200_000
#: How many workshops' stage entries are read per query while an export runs — the job never
#: holds every workshop's raw entries at once, only the flattened, masked rows it will write.
EXPORT_READ_BATCH = 50
#: How many exports one account may have waiting or running at once.
MAX_PENDING_PER_USER = 3
#: A RUNNING job older than this is assumed to belong to a worker that died, and is re-queued.
STALE_AFTER = timedelta(minutes=30)
MAX_ATTEMPTS = 3

QUEUED, RUNNING, READY, FAILED, EXPIRED = "QUEUED", "RUNNING", "READY", "FAILED", "EXPIRED"

_IDENTITY_KEY = re.compile(r"aadhaar|pehchan|cardno|identitynumber", re.IGNORECASE)


def identity_columns() -> frozenset[str]:
    """Every ``entityKey.fieldKey`` column whose value is an identity number, to be masked."""
    keys: set[str] = set()
    for stage in stages():
        for entity in stage.entities:
            for spec in entity.fields:
                fmt = getattr(spec, "text_format", None)
                if (
                    getattr(spec, "store_masked", False)
                    or fmt in (TextFormat.AADHAAR, TextFormat.PEHCHAN)
                    or _IDENTITY_KEY.search(spec.key)
                ):
                    keys.add(f"{entity.key}.{spec.key}")
    return frozenset(keys)


def mask_row(row: dict[str, str], masked: frozenset[str]) -> dict[str, str]:
    """The row with every identity cell masked. An empty cell stays empty."""
    out = dict(row)
    for key in masked.intersection(out):
        if out[key]:
            out[key] = str(mask_identity_number(out[key]) or "")
    return out


def _sheet_title(title: str, used: set[str]) -> str:
    clean = re.sub(r"[\[\]\*\?/\\:]", " ", title).strip()[:31] or "Sheet"
    candidate, n = clean, 2
    while candidate.lower() in used:
        suffix = f" ({n})"
        candidate = clean[: 31 - len(suffix)] + suffix
        n += 1
    used.add(candidate.lower())
    return candidate


class WorkbookCollector:
    """The export's rows, gathered one batch of workshops at a time, then written once.

    Split from :func:`build_workbook` so the job can read and flatten a batch, drop the raw entries,
    and read the next — and so both CPU-bound halves can run off the event loop the queue worker
    shares with the media queue and the email outbox. Pure: no database.
    """

    def __init__(self) -> None:
        self.masked = identity_columns()
        self.per_entity: dict[str, list[dict[str, str]]] = {}
        self.withheld_custom = 0
        self.unknown = 0
        self.rows_written = 0
        self.stopped = False
        self.workshops = 0

    def add(self, workshops: list[tuple[Any, list[Any]]]) -> None:
        """Flatten and mask ``[(workshop record, its live stage entries)]`` into the collector."""
        for record, entries in workshops:
            self.workshops += 1
            # No custom-section definition is passed, so a designer's own questions come back from
            # `flatten` as unreadable rows — counted as withheld, never as "newer form versions".
            grouped, unknown_entities = dw.flatten(record, entries)
            self.unknown += sum(
                item.rows for item in unknown_entities if item.entity_key != dw.CUSTOM_ENTITY_KEY
            )
            for key, rows in grouped.items():
                room = MAX_EXPORT_ROWS - self.rows_written
                if room <= 0:
                    self.stopped = True
                    break
                kept = rows[:room]
                self.stopped = self.stopped or len(kept) < len(rows)
                self.per_entity.setdefault(key, []).extend(
                    mask_row(row, self.masked) for row in kept
                )
                self.rows_written += len(kept)
            self.withheld_custom += sum(
                1 for entry in entries if getattr(entry, "entityKey", None) == dw.CUSTOM_ENTITY_KEY
            )

    def workbook(
        self,
        *,
        filters: dict[str, Any],
        scope_label: str,
        workshops_total: int,
        generated_at: datetime,
    ) -> tuple[bytes, int]:
        return _write_workbook(
            self,
            filters=filters,
            scope_label=scope_label,
            workshops_total=workshops_total,
            generated_at=generated_at,
        )


def build_workbook(
    workshops: list[tuple[Any, list[Any]]],
    *,
    filters: dict[str, Any],
    scope_label: str,
    workshops_total: int,
    generated_at: datetime,
) -> tuple[bytes, int]:
    """The .xlsx and the number of stage rows it holds. Pure: no database.

    ``workshops`` is ``[(workshop record, its live stage entries)]``. One sheet per registry entity
    that has at least one row, in registry order, each row carrying the workshop identity columns.
    """
    collector = WorkbookCollector()
    collector.add(workshops)
    return collector.workbook(
        filters=filters,
        scope_label=scope_label,
        workshops_total=workshops_total,
        generated_at=generated_at,
    )


def _write_workbook(
    collected: WorkbookCollector,
    *,
    filters: dict[str, Any],
    scope_label: str,
    workshops_total: int,
    generated_at: datetime,
) -> tuple[bytes, int]:
    from openpyxl import Workbook

    per_entity = collected.per_entity
    withheld_custom = collected.withheld_custom
    unknown = collected.unknown
    rows_written = collected.rows_written
    stopped = collected.stopped
    included = collected.workshops

    wb = Workbook(write_only=True)
    used: set[str] = set()
    readme = wb.create_sheet(_sheet_title("About this file", used))
    readme.append(["Ministry stage-data export"])
    readme.append(["Generated", generated_at.strftime("%d %b %Y %H:%M UTC")])
    readme.append(["Scope", scope_label])
    narrowing = ", ".join(f"{key}: {value}" for key, value in filters.items() if value) or "None"
    readme.append(["Filters", narrowing])
    readme.append(["Workshops", f"{included} of {workshops_total}"])
    readme.append(["Stage rows", rows_written])
    readme.append(["Identity numbers", "Masked to the last four characters."])
    readme.append(["Photographs and files", "Listed by id only."])
    if withheld_custom:
        readme.append(["Designers' own questions", f"{withheld_custom} rows not included."])
    if unknown:
        readme.append(["Newer form versions", f"{unknown} rows not shown."])
    if included < workshops_total:
        readme.append(["Stopped", f"Only the first {included} workshops are included."])
    if stopped:
        readme.append(["Stopped", f"Only the first {MAX_EXPORT_ROWS} stage rows are included."])

    identity_cols = [column.key for column in dw.WORKSHOP_IDENTITY_COLUMNS]
    identity_labels = [column.label for column in dw.WORKSHOP_IDENTITY_COLUMNS]
    for table in dw.tables():
        rows = per_entity.get(table.entity_key)
        if not rows:
            continue
        sheet = wb.create_sheet(_sheet_title(table.title, used))
        columns = [column.key for column in table.columns]
        labels = [
            f"{column.label} ({column.unit})" if column.unit else column.label
            for column in table.columns
        ]
        sheet.append([*identity_labels, "Row", *labels])
        for row in rows:
            sheet.append(
                [
                    *(row.get(key, "") for key in identity_cols),
                    row.get("entry.ordinal", ""),
                    *(row.get(key, "") for key in columns),
                ]
            )
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue(), rows_written


def export_filename(at: datetime) -> str:
    # NAMING 2.1: a snapshot that can repeat within a day carries its time.
    return naming.export_filename(
        "Ministry", "Export", "xlsx", at=at, with_time=True, subject="Stage Data"
    )


def job_payload(row: Any) -> dict[str, Any]:
    """One export as the dashboard lists it. Never the file itself.

    ``row`` is a model or a raw-query mapping (:func:`list_jobs`), so both shapes are read.
    """

    def get(key: str) -> Any:
        return row.get(key) if isinstance(row, dict) else getattr(row, key, None)

    def iso(value: Any) -> str | None:
        if isinstance(value, datetime):
            return value.isoformat()
        return value if isinstance(value, str) and value else None

    filters = get("filters")
    if isinstance(filters, str):
        try:
            filters = json.loads(filters)
        except ValueError:
            filters = {}

    return {
        "id": get("id"),
        "status": get("status"),
        "filters": filters or {},
        "workshopCount": get("workshopCount"),
        "rowCount": get("rowCount"),
        "fileName": get("fileName"),
        "fileSizeBytes": get("fileSizeBytes"),
        "error": get("error"),
        "createdAt": iso(get("createdAt")),
        "finishedAt": iso(get("finishedAt")),
        "expiresAt": iso(get("expiresAt")),
    }


#: Every column the list shows and NOT ``content``. The list is polled while a job runs, and a
#: model read would pull each stored workbook (several MB apiece) across the wire every poll only to
#: drop it; prisma-client-py's ``find_many`` has no ``select``, so the read is spelled out here.
_LIST_SQL = (
    'SELECT "id", "status", "filters", "workshopCount", "rowCount", "fileName", "fileSizeBytes", '
    '"error", "createdAt", "finishedAt", "expiresAt" FROM "MinistryDataExport" '
    'WHERE "requestedById" = $1 ORDER BY "createdAt" DESC LIMIT $2'
)


async def list_jobs(user_id: str, *, take: int = 20) -> list[dict[str, Any]]:
    """The account's own exports, newest first, without the stored files."""
    rows = await db.query_raw(_LIST_SQL, user_id, take)
    return [job_payload(row) for row in rows]


async def _claim(row: Any, worker_id: str) -> bool:
    now = datetime.now(UTC)
    claimed = await db.ministrydataexport.update_many(
        where={"id": row.id, "status": QUEUED},
        data={
            "status": RUNNING,
            "lockedAt": now,
            "lockedBy": worker_id,
            "startedAt": now,
            "attempts": {"increment": 1},
        },
    )
    return bool(claimed)


async def run_export(row: Any) -> None:
    """Build one claimed export and store the file, or record why it failed."""
    from app.api.routes.ministry_dashboard import _design_where
    from app.core.deps import can_see_ministry_dashboard
    from app.services import ministry_dashboard as register

    user = await db.user.find_unique(where={"id": row.requestedById}) if row.requestedById else None
    if user is None or not can_see_ministry_dashboard(user):
        await db.ministrydataexport.update(
            where={"id": row.id},
            data={
                "status": FAILED,
                "error": "The account that asked for this export can no longer read the dashboard.",
                "finishedAt": datetime.now(UTC),
                "lockedAt": None,
            },
        )
        return
    filters = dict(row.filters or {})
    where = _design_where(
        user,
        filters.get("search"),
        filters.get("standing"),
        filters.get("attention"),
        filters.get("state"),
    )
    if filters.get("attention") == "setupIncomplete":
        ids, _capped = await register.setup_incomplete_ids(
            _design_where(
                user, filters.get("search"), filters.get("standing"), None, filters.get("state")
            )
        )
        where.setdefault("AND", []).append({"id": {"in": ids}})
    total = await db.designworkshop.count(where=where)
    records = await db.designworkshop.find_many(
        where=where, order=[{"updatedAt": "desc"}, {"id": "asc"}], take=MAX_EXPORT_WORKSHOPS
    )
    # ONE BATCH OF WORKSHOPS AT A TIME, flattened off the event loop: the raw entries of a batch are
    # dropped before the next is read, and openpyxl never runs on the loop the worker shares.
    collector = WorkbookCollector()
    for start in range(0, len(records), EXPORT_READ_BATCH):
        batch = records[start : start + EXPORT_READ_BATCH]
        entries = await db.dwstageentry.find_many(
            where={"designWorkshopId": {"in": [str(r.id) for r in batch]}, "deletedAt": None},
            # `createdAt`, `id`: a save's new rows share one `create_many` timestamp; `id` makes it total.
            order=[{"designWorkshopId": "asc"}, {"ordinal": "asc"}, {"createdAt": "asc"}, {"id": "asc"}],
        )
        by_workshop: dict[str, list[Any]] = {}
        for entry in entries:
            by_workshop.setdefault(str(entry.designWorkshopId), []).append(entry)
        await asyncio.to_thread(
            collector.add, [(record, by_workshop.get(str(record.id), [])) for record in batch]
        )
        del entries, by_workshop
        if collector.stopped:
            break
    now = datetime.now(UTC)
    payload, rows = await asyncio.to_thread(
        collector.workbook,
        filters=filters,
        scope_label=register.scope_label(user),
        workshops_total=total,
        generated_at=now,
    )
    from prisma.fields import Base64

    # `update_many`, which answers a count: `update` would read the stored workbook straight back.
    await db.ministrydataexport.update_many(
        where={"id": row.id},
        data={
            "status": READY,
            "content": Base64.encode(payload),
            "fileName": export_filename(now),
            "fileSizeBytes": len(payload),
            "workshopCount": len(records),
            "rowCount": rows,
            "finishedAt": now,
            "expiresAt": now + timedelta(days=EXPORT_TTL_DAYS),
            "lockedAt": None,
            "error": None,
        },
    )


async def recover_and_purge(now: datetime | None = None, *, recover: bool = True) -> None:
    """Re-queue exports a dead worker left RUNNING; drop files past their keep-until date.

    ``recover=False`` skips the two stale-lock sweeps and keeps the expiry: a lock is only stale after
    :data:`STALE_AFTER` (thirty minutes), so the queue process sweeps for it once a minute rather than
    on every five-second pass (``worker.RECOVERY_INTERVAL_SECONDS``), while an expired file still stops
    being READY on the pass it expires, as before.
    """
    now = now or datetime.now(UTC)
    if recover:
        await _recover_stale(now)
    await db.ministrydataexport.update_many(
        where={"status": READY, "expiresAt": {"lt": now}},
        data={"content": None, "status": EXPIRED},
    )


async def _recover_stale(now: datetime) -> None:
    await db.ministrydataexport.update_many(
        where={
            "status": RUNNING,
            "lockedAt": {"lt": now - STALE_AFTER},
            "attempts": {"lt": MAX_ATTEMPTS},
        },
        data={"status": QUEUED, "lockedAt": None, "lockedBy": None},
    )
    await db.ministrydataexport.update_many(
        where={
            "status": RUNNING,
            "lockedAt": {"lt": now - STALE_AFTER},
            "attempts": {"gte": MAX_ATTEMPTS},
        },
        data={
            "status": FAILED,
            "error": "The export did not finish. Ask for it again.",
            "lockedAt": None,
        },
    )


async def process_next_ministry_exports(
    *, worker_id: str, limit: int = 1, recover: bool = True
) -> int:
    """The worker's turn: recover, purge, then build up to ``limit`` queued exports. Returns built.

    ``recover`` is :func:`recover_and_purge`'s: the queue process passes False between its once-a-
    minute sweeps."""
    await recover_and_purge(recover=recover)
    queued = await db.ministrydataexport.find_many(
        where={"status": QUEUED}, order={"createdAt": "asc"}, take=limit
    )
    built = 0
    for row in queued:
        if not await _claim(row, worker_id):
            continue
        try:
            await run_export(row)
            built += 1
        except Exception as exc:  # one bad export must never stop the worker
            log.exception("ministry export %s failed", row.id)
            await db.ministrydataexport.update(
                where={"id": row.id},
                data={
                    "status": FAILED,
                    "error": f"The export could not be built ({type(exc).__name__}). Ask for it again.",
                    "finishedAt": datetime.now(UTC),
                    "lockedAt": None,
                },
            )
    return built
