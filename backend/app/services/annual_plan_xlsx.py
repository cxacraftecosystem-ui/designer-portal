"""The ministry's annual directory of planned workshops: the pro-forma it is typed into, and the
reader that gets it back out.

THE THIRD .XLSX READER IN THIS REPOSITORY, and it is built on ``app/services/xlsx_table.py`` rather
than beside ``questionnaire_xlsx`` or ``artisan_xlsx``. Every primitive that is about SPREADSHEETS —
header discovery, NFKC folding, the two-pass formula map, the magic-byte refusals, ``ParseProblem``
— is shared. What lives here is the part that is about THIS DOCUMENT: which columns a ministry
directory has, how a date written by a person in an Indian office is read, and what the natural key
is.

The owner's source document marks this requirement “Phase 2”, which is a quotation out of their
paper and not a claim about this repository. What is true as of 2026-09-13: the directory uploads,
lists, exports and promotes; there is no Android screen and no per-region scope — check
``grep -rn "annual-plan" android/ frontend/lib/permissions.ts``.

══ WHAT THE PARSER OWES A PERSON WHO UPLOADS 300 ROWS ═══════════════════════════════════════════

NOTHING IS EVER SILENTLY DROPPED. A row this cannot read is reported with its Excel gutter row
number and the cell text that defeated it, and the other 299 are still imported. That is
``xlsx_table``'s contract and it matters more here, not less: a designer who loses two questions out
of forty notices at the next sitting; an administrator who loses two workshops out of three hundred
finds out when a district asks why nobody came.

A COLUMN THIS DOES NOT RECOGNISE IS REPORTED TOO, and that is the one place this parser is stricter
than the questionnaire's. A questionnaire pro-forma invites a designer to add columns of their own —
its instructions sheet says so — so an unknown heading there is normal. A ministry directory's
columns are all data, and one this reader skipped is a column somebody filled in and nobody stored,
with no symptom anywhere.

══ THE DATE IS THE DANGEROUS COLUMN, AND IT IS DANGEROUS IN A WAY NOTHING ELSE HERE IS ═══════════

``03/04/2026`` is the third of April in every office that will send this file and the fourth of
March to ``datetime.strptime`` with the American ordering. Both readings are valid dates; neither
raises; the row imports; and a workshop is scheduled a month out with no error anywhere in the
system. So:

* A REAL Excel date cell is taken exactly as openpyxl resolved it, and no rule of ours is applied —
  Excel has already decided, and second-guessing it is how a correctly-typed sheet gets corrupted.
  That is the whole reason this module reads RAW VALUES (``xlsx_table.sheet_value_rows``) rather
  than the text rows every other reader here takes: ``cell_text`` renders a date cell
  "2026-03-12 00:00:00" and the resolved value is then gone.
* A TEXT date is read DAY FIRST, because this directory is Indian.
* A TEXT date where BOTH readings are valid — every day-of-month at or below twelve — is read day
  first AND REPORTED as a warning naming both readings, so the one case where we could be silently
  wrong is the one case the administrator is shown.

That last rule produces warnings on an ordinary file, and it is meant to. The alternative is a
parser that is quietly wrong about roughly two fifths of the dates in any sheet typed as text.

══ EVERY CLIP CAP HERE MIRRORS A STAGE-1 FieldSpec, AND THE PAIR IS PINNED BY A TEST ═════════════

See ``_TEXT_CAPS`` and ``_CAP_MIRRORS``. This is not tidiness: a value longer than the registry's
bound is DROPPED AND ONLY LOGGED by ``validate_entry(enforce_required=False)`` on the promotion
path, so an over-length venue would be written as a ``DesignWorkshop`` COLUMN with no stage entry
behind it — and a column with no entry behind it is nulled by the designer's first stage-1 save
under a 200 reading "Stage saved".
"""

from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from app.services import address
from app.services.stage_schema import ENUMS
from app.services.xlsx_report import _put
from app.services.xlsx_table import (
    FormulaCell,
    ParseProblem,
    XlsxTableError,
    as_text_rows,
    cell_text,
    clip,
    load_value_sheets,
    norm_header,
    pick_sheet,
)

__all__ = [
    # Exported because the refusal that reads a declared filter has to bound it before it lands in
    # an HTTP ``detail`` string, and the bound belongs beside the writer that applies it.
    "MAX_FILTER_NOTE_CHARS",
    "MAX_PLAN_ROWS",
    "MAX_PLAN_YEAR",
    "MIN_PLAN_YEAR",
    "PRO_FORMA_FILENAME",
    "AnnualPlanXlsxError",
    "ParsedAnnualPlan",
    "ParsedPlanRow",
    "build_annual_plan_pro_forma",
    "build_annual_plan_workbook",
    "export_filename",
    "fold_workshop_no",
    "parse_annual_plan_workbook",
    "plan_year_label",
    "read_plan_date",
    "read_plan_year",
]


PRO_FORMA_FILENAME = "annual-plan-pro-forma.xlsx"

SHEET_PLAN = "Annual plan"
SHEET_DETAILS = "Details"
SHEET_HELP = "How to fill this in"

# Blue, and deliberately neither the questionnaire's purple (5B21B6) nor the artisan list's teal
# (0F766E). Three pro-formas now leave this building and an administrator with two of them open has
# only the colour of the header band to tell them apart at a glance.
_BRAND = "1D4ED8"
_WHITE = "FFFFFFFF"
_GREY = "FF6B7280"

# Ceilings. The requirement says 200-300 rows a year. These are sized so a hostile or malformed file
# cannot turn one upload into an unbounded read, and so the write in `annual_plan.apply_parsed_plan`
# stays one honest transaction: 600 rows bounds the worst case at one create_many plus 600 updates
# plus one update_many, inside the 60 s transaction timeout that call opens.
MAX_PLAN_ROWS = 600
MAX_ROWS_SCANNED = 5_000  # a directory with a title block, blank spacer rows and a total row
MAX_SHEETS_SCANNED = 12
MAX_HEADER_SCAN_ROWS = 15

# ── THE CLIP CAPS, AND WHY EACH NUMBER IS THE NUMBER IT IS ───────────────────────────────────────
#
# EVERY ONE OF THESE MIRRORS A stage-1 `workshopSetup` FieldSpec's `max_length`, and it is the
# FieldSpec that decides, not this file. `annual_plan.promote_entry` hands these values to
# `seed_designer_prefill`, which runs `validate_entry(..., enforce_required=False)` — and a value
# that validator REFUSES is dropped and only logged (see its own comment block in
# `services/design_workshops.py`, "a value refused here vanished completely"). The promoted COLUMN
# is still written. A column with no stage entry behind it is exactly what the designer's first
# stage-1 save nulls out, under a 200 reading "Stage saved", with nothing warning and completeness
# unmoved. So a cap here that is one character looser than its FieldSpec is a data-loss bug that
# only appears at the boundary and only days later.
#
# `_CAP_MIRRORS` below names the pair for `tests/test_annual_plan_promotion.py::
# test_no_parser_clip_cap_exceeds_the_field_it_is_promoted_into`, which fails if either side moves.
_TEXT_CAPS: dict[str, int] = {
    "workshopNo": 60,  # workshopSetup.workshopCode
    "plannedTitle": 220,  # workshopSetup.workshopTitle, and DesignWorkshopCreate.title
    "craftName": 160,  # workshopSetup.craftName
    "clusterName": 160,  # workshopSetup.clusterName
    "state": 80,  # workshopSetup.state
    "district": 80,  # workshopSetup.district
    "venue": 220,  # workshopSetup.venue
    "implementingAgency": 220,  # workshopSetup.implementingAgency
    "sponsor": 220,  # workshopSetup.sponsor
    # NOT a promoted column and not in `_CAP_MIRRORS`: remarks stay in the directory and are never
    # seeded into a stage, so this number answers to legibility rather than to the registry.
    "notes": 2_000,
}

#: The parser role → the ``workshopSetup`` field key its value is promoted into. Consumed only by
#: the contract test; kept beside the caps so the two cannot be edited apart.
_CAP_MIRRORS: dict[str, str] = {
    "workshopNo": "workshopCode",
    "plannedTitle": "workshopTitle",
    "craftName": "craftName",
    "clusterName": "clusterName",
    "state": "state",
    "district": "district",
    "venue": "venue",
    "implementingAgency": "implementingAgency",
    "sponsor": "sponsor",
}

# The plan year a directory may name. Below the first, somebody has typed a row number into the year
# box; above the last, a typo. Both are refused with the number in the sentence, because "that is
# not a valid year" about a cell the person cannot see is a message they cannot act on.
MIN_PLAN_YEAR = 2000
MAX_PLAN_YEAR = 2100


class AnnualPlanXlsxError(ValueError):
    """The upload is not a directory this can read at all — the whole file, not one row.

    Distinct from a :class:`ParseProblem`, exactly as ``ArtisanXlsxError`` and
    ``QuestionnaireXlsxError`` are. The message is written to be shown to the administrator verbatim
    and reaches them as an HTTP 422 ``detail``.
    """


# --- The column vocabulary ----------------------------------------------------------------------

# Roles are tried in this dict's ORDER, so a more specific spelling is matched before a looser one
# could swallow it — the rule `xlsx_table._role_for` states about `aliases` being tried in its own
# order.
#
# ── WHY "Sl. No." HAS ITS OWN ROLE AND IS THEN THROWN AWAY ───────────────────────────────────────
#
# It is the single most dangerous heading in this document. Almost every ministry directory carries
# a serial number in column A — 1, 2, 3 — beside the real workshop reference in column B. Aliasing
# "sl no" to `workshopNo` "because it is a number that identifies the row" keys three hundred rows
# on 1..300, and then a corrected sheet with ONE workshop inserted at the top REWRITES EVERY ROW
# BELOW IT with its neighbour's data, silently, under a 201 saying the upload succeeded. There is no
# constraint violation, no problem reported, and nothing on any screen that could show it — the row
# count is right, every field is populated, and the venues have all moved down one.
#
# So the serial number is RECOGNISED — which is what stops the unknown-column warning firing on a
# perfectly ordinary sheet, and what stops it being mistaken for anything else — and its value is
# discarded. A sheet that carries ONLY a serial number and no workshop reference is refused WHOLE,
# naming the heading it needs, because such a sheet has no stable key and no correction to it could
# ever be applied.
_COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "serialNo": ("sl no", "sl", "s no", "sr no", "serial no", "serial number", "sno", "srno", "no"),
    "workshopNo": (
        "workshop no",
        "workshop number",
        "workshop code",
        "workshop id",
        "workshop ref",
        "workshop reference",
        "ws no",
        "wsno",
        "sanction no",
        "sanction number",
    ),
    "plannedTitle": ("workshop title", "name of workshop", "workshop name", "title", "name"),
    "workshopKind": (
        "workshop kind",
        "type of workshop",
        "workshop type",
        "kind",
        "type",
        "category",
    ),
    "craftName": ("craft", "craft name", "handicraft", "craft handicraft"),
    "clusterName": ("cluster", "cluster name", "block"),
    "state": ("state", "state ut", "state union territory", "state ut name"),
    "district": ("district", "district name"),
    "venue": ("venue", "venue name", "venue address", "place", "place of workshop", "location"),
    "plannedStartDate": (
        "date",
        "start date",
        "from date",
        "date of workshop",
        "workshop date",
        "planned date",
        "commencement date",
        "date from",
        "start",
    ),
    "plannedEndDate": ("end date", "to date", "till date", "date to", "closing date", "end"),
    "implementingAgency": (
        "implementing agency",
        "agency",
        "implementing body",
        "implementing organisation",
        "implementing organization",
    ),
    "sponsor": ("sponsor", "sponsoring body", "funded by", "funding body"),
    "notes": ("notes", "note", "remarks", "remark", "comment", "comments", "observation"),
}

_DETAIL_KEYS: dict[str, tuple[str, ...]] = {
    # TWO KEYS, AND THE SECOND ONE IS NOT A ROW ADDED FOR A READER'S BENEFIT. Every label on the
    # Details sheet is matched against these aliases, so a row added here for decoration ("Issued
    # by", "Uploaded on") becomes a value a parser STORES — `xlsx_table.read_details`' own warning
    # block records what that cost the questionnaire, where a row labelled "Notes" silently became
    # the questionnaire's description on re-import. Both keys below are read into a FLAG or a YEAR
    # on `ParsedAnnualPlan` and neither can reach a plan column: `_COLUMN_ALIASES` is a separate
    # table and the two share no name.
    "planYear": ("plan year", "financial year", "year", "fy", "annual plan year", "plan"),
    # ── WHY A WORKBOOK HAS TO BE ABLE TO SAY "I AM ONLY PART OF THE YEAR" ─────────────────────────
    #
    # `GET /annual-plan/export.xlsx` applies the LIST's filters, and `POST /annual-plan/upload` with
    # "withdraw the workshops missing from this sheet" ticked treats every row of the year that the
    # sheet does not name as absent and stamps `withdrawnAt` on it. Those two are safe apart and
    # ruinous together: filter the 2026-27 list to forty rows, press Export, correct two venues,
    # re-upload with the box ticked, and the other two hundred and sixty are withdrawn in one
    # `update_many`. The export route's own docstring names this hazard — "a page of fifty would be
    # a corrected sheet that silently withdrew the other two hundred and fifty on re-upload if they
    # ticked the box" — and closes it for PAGING while leaving the filters producing exactly that
    # shape, under a filename (`annual-plan-2026-27.xlsx`) that reads as the whole year.
    #
    # The workbook is the only place the fact survives the round trip: the upload sees bytes and a
    # form field, never the query string the download was taken under. So a filtered export WRITES
    # this row, `_details_filter` reads it back, and the upload route refuses `withdrawAbsent` for a
    # sheet that declares one. An unfiltered export writes nothing, so the ordinary file is
    # unchanged byte for byte.
    "filter": ("filtered view", "filtered export", "filter", "partial view", "partial export"),
}

#: How much of a declared filter is carried off a Details sheet into a refusal message.
#:
#: BOUNDED BECAUSE IT IS UNTRUSTED TEXT FROM AN UPLOADED FILE and it lands in an HTTP ``detail``
#: string. Nobody who exported a filtered view has a filter sentence anywhere near this long; a
#: value that is, is somebody's paragraph typed into the wrong cell, and the refusal needs to stay
#: readable rather than faithful.
MAX_FILTER_NOTE_CHARS = 200

_DETAIL_SHEET_NAMES = ("details", "about", "plan details", "info", "information")

_WIDTHS: dict[str, int] = {
    "workshopNo": 20,
    "plannedTitle": 40,
    "workshopKind": 26,
    "craftName": 22,
    "clusterName": 22,
    "state": 22,
    "district": 22,
    "venue": 40,
    "plannedStartDate": 14,
    "plannedEndDate": 14,
    "implementingAgency": 28,
    "sponsor": 24,
    "notes": 34,
}

_PRO_FORMA_COLUMNS: tuple[tuple[str, str], ...] = (
    ("workshopNo", "Workshop No."),
    ("plannedTitle", "Workshop Title"),
    ("workshopKind", "Type of Workshop"),
    ("craftName", "Craft"),
    ("clusterName", "Cluster"),
    ("state", "State"),
    ("district", "District"),
    ("venue", "Venue"),
    ("plannedStartDate", "Start Date"),
    ("plannedEndDate", "End Date"),
    ("implementingAgency", "Implementing Agency"),
    ("sponsor", "Sponsoring Body"),
    ("notes", "Remarks"),
)

#: The one column an administrator MUST fill in, coloured differently on the header band.
_REQUIRED = frozenset({"workshopNo"})

#: Roles whose value is actually stored on a row. `serialNo` is recognised and discarded — see the
#: block above `_COLUMN_ALIASES`.
_STORED_ROLES = frozenset(_COLUMN_ALIASES) - {"serialNo"}


# --- What comes out ------------------------------------------------------------------------------


@dataclass
class ParsedPlanRow:
    """One planned workshop, as the sheet spells it. NOT a workshop — see ``AnnualPlanEntry``."""

    sheetRow: int
    workshopNo: str
    workshopNoKey: str
    plannedTitle: str | None = None
    workshopKind: str | None = None
    craftName: str | None = None
    clusterName: str | None = None
    state: str | None = None
    district: str | None = None
    venue: str | None = None
    plannedStartDate: date | None = None
    plannedEndDate: date | None = None
    implementingAgency: str | None = None
    sponsor: str | None = None
    notes: str | None = None


@dataclass
class ParsedAnnualPlan:
    planYear: int | None = None
    sheet: str | None = None
    #: The filter the workbook DECLARES it was exported under, verbatim, or ``None`` for a sheet
    #: that declares none. It is a SENTENCE rather than a flag because the only consumer is a
    #: refusal shown to an administrator, and "this sheet was exported with Search: 'Bhuj'" is
    #: actionable where "partial: true" is not. See :data:`_DETAIL_KEYS` for what it is for.
    filterNote: str | None = None
    rows: list[ParsedPlanRow] = field(default_factory=list)
    problems: list[ParseProblem] = field(default_factory=list)

    @property
    def rowCount(self) -> int:
        return len(self.rows)

    def payload(self) -> dict[str, Any]:
        return {
            "planYear": self.planYear,
            "sheet": self.sheet,
            "filterNote": self.filterNote,
            "rowCount": self.rowCount,
            "problems": [p.payload() for p in self.problems],
        }


# --- The key fold ---------------------------------------------------------------------------------

#: Whitespace of every kind Excel and Word between them can put in a cell, including the
#: non-breaking space a reference pasted out of a PDF routinely carries.
#:
#: EVERY CHARACTER IS SPELLED AS AN ESCAPE. Written as literals they are invisible in a diff
#: and in a review, and one of them (U+200B, the zero-width space) is invisible in an editor
#: too — which is precisely the property that makes it worth folding away in the first place.
_SPACEY = re.compile(r"[\s\u00a0\u1680\u2000-\u200b\u202f\u205f\u3000]+")


def fold_workshop_no(value: Any) -> str:
    """The matching form of a ministry workshop reference.

    THE ONLY WRITER OF ``AnnualPlanEntry.workshopNoKey``, and there is deliberately no SQL twin of
    it — the migration's header says why. NFKC first, because a reference pasted out of a PDF
    routinely carries full-width digits and a non-breaking space; then every run of whitespace
    collapsed to one ordinary space; then trimmed and upper-cased.

    PUNCTUATION IS KEPT. "DPW/2026/017" and "DPW-2026-017" stay two different references, and that
    is right: they are two different strings in the ministry's own document, and folding them
    together would merge two plan rows on a guess. The folding this does is confined to spellings
    that are INVISIBLE on screen — the ones a person cannot see they have typed differently.
    """
    text = unicodedata.normalize("NFKC", str(value or ""))
    return _SPACEY.sub(" ", text).strip().upper()


def plan_year_label(year: int | None) -> str:
    """2026 -> "2026-27". The printable form of the financial year, rendered and never stored.

    Stored as one integer because an integer sorts, indexes and compares with no normalisation rule,
    and a normalisation rule spelled in Python and again in SQL is the trap the identity keys carry.
    The label is built here so that the server owns the words and the two clients cannot spell one
    year two ways.
    """
    if year is None:
        return "—"
    return f"{year}-{(year + 1) % 100:02d}"


def export_filename(year: int | None) -> str:
    """``annual-plan-2026-27.xlsx`` — pure ASCII, which is load-bearing for the download header.

    ``xlsx_report.xlsx_response`` writes a plain ``filename="…"`` content-disposition. The RFC 6266
    ``filename*`` form that ``data_browser._content_disposition`` builds exists for Devanagari RECORD
    names and is not needed here, because every character this function can emit is an ASCII digit
    or a hyphen. Said out loud so nobody "fixes" the simpler header into the harder one.
    """
    return f"annual-plan-{plan_year_label(year)}.xlsx" if year is not None else "annual-plan.xlsx"


# --- Reading a plan year --------------------------------------------------------------------------

#: "2026", "2026-27", "2026-2027", "2026/27", with an optional FY prefix and optional punctuation.
_YEAR_PATTERNS = (
    # `\d{1,4}` AND NOT `\d{4}`, so "17" — a row number typed into the year box, which is the
    # commonest way this cell goes wrong — reaches the MIN/MAX refusal that names the bound, rather
    # than the generic "not a year this can read" that names nothing the administrator can act on.
    re.compile(r"^(?P<a>\d{1,4})$"),
    re.compile(r"^(?P<a>\d{4})\s*[-/–—]\s*(?P<b>\d{2})$"),
    re.compile(r"^(?P<a>\d{4})\s*[-/–—]\s*(?P<b>\d{4})$"),
)
#: Everything an office writes in front of the year: "FY", "F.Y.", "Financial Year", "Year".
_YEAR_PREFIX = re.compile(
    r"^(?:f\s*\.?\s*y\s*\.?|financial\s+year|fin\.?\s*year|year|plan\s+year)\s*[:\-]?\s*",
    re.IGNORECASE,
)


def read_plan_year(
    value: Any, *, sheet: str, row: int | None, problems: list[ParseProblem]
) -> int | None:
    """The financial year a Details sheet names, as its STARTING calendar year, or ``None``.

    SEVEN SPELLINGS AND THREE CELL TYPES, because this is one cell on a sheet nobody proof-reads and
    getting it wrong files three hundred rows under a year the document does not claim. A person
    types "2026-27"; a person in a hurry types "2026"; a person copying the ministry's own header
    types "F.Y. 2026-2027"; and Excel, left to itself, turns "2026-27" into the DATE 27 June 2026 or
    stores a bare "2026" as the number 2026. All of those are answered; anything else is refused
    with the text in the sentence so the administrator can find the cell.

    A YEAR OUTSIDE ``MIN_PLAN_YEAR``..``MAX_PLAN_YEAR`` IS AN ERROR AND NOT A CLAMP. Below the
    floor somebody has typed a row number into the year box; above the ceiling it is a typo. Either
    way the honest answer is "I do not know which year this is", because filing a directory under a
    guessed year is the one mistake a re-upload cannot repair — the corrected sheet lands in the
    right year and the three hundred wrong rows stay where they are.
    """
    if value is None:
        return None
    if isinstance(value, FormulaCell):
        problems.append(
            ParseProblem(
                sheet,
                row,
                "warning",
                (
                    "The plan year cell holds a formula whose result is not saved in the file, so "
                    "the year could not be read. Open the workbook in Excel, press F9, save, and "
                    "upload again — or choose the year on the upload panel."
                ),
                str(value),
            )
        )
        return None
    # EXCEL TURNED IT INTO A DATE, WHICH IS THE COMMONEST WAY THIS CELL GOES WRONG. "2026-27" typed
    # into a General cell is auto-formatted as 27 June 2026 by every recent Excel. The year is still
    # recoverable and the row is still usable, so it is taken — and said out loud, because the
    # administrator's workbook now shows a date where they typed a financial year and they will meet
    # that again next year.
    if isinstance(value, datetime | date):
        year = value.year
        if MIN_PLAN_YEAR <= year <= MAX_PLAN_YEAR:
            problems.append(
                ParseProblem(
                    sheet,
                    row,
                    "warning",
                    (
                        f"Excel stored the plan year as a date, so it was read as {year} "
                        f"({plan_year_label(year)}). Format that cell as Text and type the year "
                        "again if that is not the year you meant."
                    ),
                    value.isoformat(),
                )
            )
            return year
        return _refuse_year(str(year), sheet, row, problems)
    if isinstance(value, int | float) and not isinstance(value, bool):
        year = int(value)
        if MIN_PLAN_YEAR <= year <= MAX_PLAN_YEAR:
            return year
        return _refuse_year(str(year), sheet, row, problems)

    text = cell_text(value)
    if not text:
        return None
    stripped = _YEAR_PREFIX.sub("", text.replace(".", " ").strip()).strip()
    stripped = _SPACEY.sub(" ", stripped).strip()
    for pattern in _YEAR_PATTERNS:
        match = pattern.match(stripped)
        if not match:
            continue
        start = int(match.group("a"))
        second = match.groupdict().get("b")
        if second is not None:
            # "2026-27" and "2026-2027" both mean one financial year. "2026-29" does not mean
            # anything this table can hold, and quietly keeping the first half would file a
            # three-year programme as a one-year plan.
            expected_short = (start + 1) % 100
            if (len(second) == 2 and int(second) != expected_short) or (
                len(second) == 4 and int(second) != start + 1
            ):
                problems.append(
                    ParseProblem(
                        sheet,
                        row,
                        "error",
                        (
                            f"'{text}' is not one financial year. A financial year runs from one "
                            f"April to the next, so it is written {plan_year_label(start)} or "
                            f"{start}-{start + 1}. Correct the 'Plan year' cell and upload again."
                        ),
                        text,
                    )
                )
                return None
        if MIN_PLAN_YEAR <= start <= MAX_PLAN_YEAR:
            return start
        return _refuse_year(text, sheet, row, problems)

    problems.append(
        ParseProblem(
            sheet,
            row,
            "error",
            (
                f"'{text}' is not a year this can read. Type the financial year as "
                f"{plan_year_label(2026)} or 2026, or choose the year on the upload panel."
            ),
            text,
        )
    )
    return None


def _refuse_year(text: str, sheet: str, row: int | None, problems: list[ParseProblem]) -> None:
    problems.append(
        ParseProblem(
            sheet,
            row,
            "error",
            (
                f"'{text}' is not a plan year this directory can hold — the year has to be between "
                f"{MIN_PLAN_YEAR} and {MAX_PLAN_YEAR}. Correct the 'Plan year' cell and upload "
                "again."
            ),
            text,
        )
    )


# --- Reading a date -------------------------------------------------------------------------------

_TEXT_DATE_ORDERS: tuple[tuple[str, bool], ...] = (
    # DAY FIRST, ALWAYS FIRST. See the module docstring. The boolean says "this pattern is ambiguous
    # with its month-first twin", which is what triggers the warning.
    ("%d/%m/%Y", True),
    ("%d-%m-%Y", True),
    ("%d.%m.%Y", True),
    ("%d/%m/%y", True),
    ("%d-%m-%y", True),
    ("%Y-%m-%d", False),
    ("%Y/%m/%d", False),
    ("%d %b %Y", False),
    ("%d %B %Y", False),
    ("%b %d %Y", False),
    ("%B %d %Y", False),
    ("%d-%b-%Y", False),
    ("%d-%B-%Y", False),
)


def read_plan_date(
    value: Any, *, sheet: str, row: int, column_label: str, problems: list[ParseProblem]
) -> date | None:
    """One date cell. Returns the date, or ``None`` having reported why.

    A real Excel date arrives as ``datetime``/``date`` and is taken AS-IS: Excel has already
    resolved its serial number and re-reading its string form would be the corruption this function
    exists to prevent.

    ⚠ THIS RETURNS ``None`` AFTER AN ``"error"`` AND THE ROW IS STILL IMPORTED. The severity says
    THIS CELL was not stored, not that the row was dropped — a planned workshop with an unreadable
    date is still a workshop the ministry planned, and refusing the row would lose the district and
    the venue with it. This is the ONE place in this module where ``"error"`` does not mean a lost
    row, and a reader must not generalise from it: everywhere else an ``"error"`` is a row that did
    not import.
    """
    if isinstance(value, FormulaCell):
        problems.append(
            ParseProblem(
                sheet,
                row,
                "warning",
                (
                    f"The {column_label} cell holds a formula whose result is not saved in the "
                    "file, so the date could not be read. Open the workbook in Excel, press F9, "
                    "save, and upload again."
                ),
                str(value),
            )
        )
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = cell_text(value)
    if not text:
        return None
    for pattern, ambiguous_order in _TEXT_DATE_ORDERS:
        try:
            # NAIVE ON PURPOSE (ruff DTZ007). This is a calendar DAY off a spreadsheet and not an
            # instant: a workshop planned for 12 March is planned for 12 March in Bhuj and in Delhi
            # alike, and attaching a zone here would move it by a day for half the country the
            # first time somebody read it back in another one.
            parsed = datetime.strptime(text, pattern).date()  # noqa: DTZ007
        except ValueError:
            continue
        if ambiguous_order and parsed.day <= 12 and parsed.month != parsed.day:
            other = parsed.replace(day=parsed.month, month=parsed.day)
            problems.append(
                ParseProblem(
                    sheet,
                    row,
                    "warning",
                    (
                        f"The {column_label} '{text}' was read as {parsed.isoformat()} (day "
                        f"first). If you meant {other.isoformat()}, format that column as a date "
                        "in Excel and upload again — a date cell is never guessed at."
                    ),
                    text,
                )
            )
        return parsed
    problems.append(
        ParseProblem(
            sheet,
            row,
            "error",
            (
                f"The {column_label} '{text}' is not a date this can read. Type it as 12/03/2026, "
                "2026-03-12 or 12 March 2026 — or format the column as a date in Excel."
            ),
            text,
        )
    )
    return None


# --- The workshop kind ----------------------------------------------------------------------------


def _read_workshop_kind(
    text: str, *, sheet: str, row: int, problems: list[ParseProblem]
) -> str | None:
    """A registry token, or ``None`` having said which six words are allowed.

    NEVER STORED RAW. `AnnualPlanEntry.workshopKind` holds a token from
    `stage_schema.ENUMS["WORKSHOP_KIND"]` and nothing else, because promotion copies it onto
    `DesignWorkshop.workshopKind`, which the workshop list narrows by and the report cover prints.
    A kind spelled in the ministry's own words that the registry has never heard of would reach a
    dropdown that cannot show it and a cover page that would print it.
    """
    kinds: dict[str, str] = ENUMS.get("WORKSHOP_KIND", {})
    folded = norm_header(text)
    for token, label in kinds.items():
        if folded in (norm_header(token), norm_header(label)):
            return token
    problems.append(
        ParseProblem(
            sheet,
            row,
            "warning",
            (
                f"'{text}' is not a type of workshop this app knows, so the type was left blank "
                f"and everything else on the row was kept. Use one of: "
                f"{', '.join(kinds.values())}."
            ),
            text,
        )
    )
    return None


# --- Reading the workbook --------------------------------------------------------------------------


def parse_annual_plan_workbook(data: bytes, *, filename: str | None = None) -> ParsedAnnualPlan:
    """Read an uploaded directory into rows.

    Raises :class:`AnnualPlanXlsxError` ONLY when the FILE is unusable; anything wrong with a row of
    an otherwise-readable file comes back in ``result.problems`` and the other rows are still
    offered. An administrator with 298 of 300 workshops and a list of the two that need correcting
    is far better off than one with an error page.
    """
    if not data:
        raise AnnualPlanXlsxError(
            "The upload was empty. Attach the filled-in annual plan pro-forma."
        )
    try:
        # VALUE ROWS, NOT TEXT ROWS, AND THAT IS THE WHOLE REASON THIS READER DIFFERS FROM THE OTHER
        # TWO — see the module docstring's date paragraph and `xlsx_table.sheet_value_rows`.
        sheets = load_value_sheets(
            data,
            max_sheets=MAX_SHEETS_SCANNED,
            max_rows=MAX_ROWS_SCANNED,
            pro_forma="the annual plan pro-forma",
        )
    except XlsxTableError as exc:
        raise AnnualPlanXlsxError(str(exc)) from exc

    if not sheets:
        raise AnnualPlanXlsxError("The workbook has no sheets in it.")

    result = ParsedAnnualPlan()

    # THE HEADER SCAN RUNS OVER TEXT ROWS AND THE DATA READ RUNS OVER VALUE ROWS, deliberately.
    # `pick_sheet`/`scan_header` declare `list[tuple[int, list[str]]]` and mean it; handing them raw
    # values would work today by accident (`norm_header` stringifies anything) and would be a lie in
    # the signature that the next reader has to disprove.
    text_sheets = {name: as_text_rows(rows) for name, rows in sheets.items()}

    # The Details sheet, if there is one. Missing is normal — an office that built the directory
    # from scratch has none, and the route can supply the year instead.
    for name, rows in sheets.items():
        if norm_header(name) in _DETAIL_SHEET_NAMES:
            result.planYear = _details_plan_year(rows, name, result.problems)
            result.filterNote = _details_filter(rows)
            break

    picked = pick_sheet(
        text_sheets,
        result.problems,
        aliases=_COLUMN_ALIASES,
        required_role="workshopNo",
        preferred_names=(SHEET_PLAN, "annual plan", "plan", "workshops", "sheet1"),
        preferred_display=SHEET_PLAN,
        looks_like="an annual plan",
        most_noun="rows",
        max_scan_rows=MAX_HEADER_SCAN_ROWS,
    )
    if picked is None:
        raise AnnualPlanXlsxError(
            "No sheet in that workbook has a 'Workshop No.' column, so there was nothing to "
            "import. Download the pro-forma and type your directory under its headings, or add a "
            "header row with a column called 'Workshop No.'. A sheet with only a serial number "
            "cannot be used: a corrected sheet has to be able to find the rows it already wrote, "
            "and 1, 2, 3 changes the moment a row is inserted."
        )
    sheet_name, header_row, columns = picked
    result.sheet = sheet_name

    _report_unrecognised_columns(
        text_sheets[sheet_name], sheet_name, header_row, columns, result.problems
    )
    _read_rows(sheets[sheet_name], sheet_name, header_row, columns, result)

    # THE FILENAME IS THE LAST FALLBACK AND THERE IS NO FURTHER ONE HERE. The route supplies the
    # third source (the administrator's own year picker) and refuses the upload when all three are
    # silent, because a directory filed under a guessed year is the mistake a re-upload cannot fix.
    if result.planYear is None and filename:
        match = re.search(r"(20\d\d)", filename)
        if match:
            year = int(match.group(1))
            if MIN_PLAN_YEAR <= year <= MAX_PLAN_YEAR:
                result.planYear = year

    if not result.rows:
        raise AnnualPlanXlsxError(
            f"The '{sheet_name}' sheet has a Workshop No. column but no rows under it. Type at "
            "least one planned workshop and upload again."
        )
    return result


def _report_unrecognised_columns(
    rows: list[tuple[int, list[str]]],
    sheet: str,
    header_row: int,
    columns: dict[int, tuple[str, str | None]],
    problems: list[ParseProblem],
) -> None:
    """ONE warning naming every heading on the header row this reader did not recognise.

    ⚠ ONLY HEADINGS THAT HAVE TEXT IN THEM. openpyxl pads every row out to ``ws.max_column``, so a
    sheet with stray formatting out to column BZ has forty header cells holding ``None`` — and
    `norm_header(None)` is `""`, which matches no alias. Without the non-empty filter below, EVERY
    ORDINARY UPLOAD produces a warning reading "These columns were not recognised and nothing in
    them was stored: , , , , , …", which teaches an administrator to ignore this panel, which is the
    panel that tells them a column of real data was skipped.

    This is computed HERE rather than returned by `xlsx_table.scan_header`, on purpose: that
    function is shared with the questionnaire and the artisan list, both of which INVITE columns of
    the uploader's own (their instructions sheets say so) and neither of which wants this warning. A
    third element on a shared return type that two of its three callers must discard is a signature
    that exists for one caller.
    """
    header = next((values for number, values in rows if number == header_row), [])
    unknown = [
        text
        for index, raw in enumerate(header)
        if index not in columns and (text := cell_text(raw).strip())
    ]
    if not unknown:
        return
    problems.append(
        ParseProblem(
            sheet,
            header_row,
            "warning",
            (
                "These columns were not recognised and nothing in them was stored: "
                f"{', '.join(unknown)}. Rename them to one of the pro-forma's headings if the "
                "information should be kept."
            ),
            None,
        )
    )


def _clip_role(
    value: str | None, role: str, sheet: str, row: int, problems: list[ParseProblem]
) -> str | None:
    """Hold one stored value to its `_TEXT_CAPS` bound, reporting the clip. Never silent.

    A clip that says nothing is a ministry's venue quietly losing its last forty characters. A clip
    that does not happen at all is worse — see the `_TEXT_CAPS` block above for what an over-length
    value does on the promotion path.
    """
    if value is None:
        return None
    cap = _TEXT_CAPS[role]
    if len(value) <= cap:
        return value
    label = dict(_PRO_FORMA_COLUMNS)[role]
    problems.append(
        ParseProblem(
            sheet,
            row,
            "warning",
            (
                f"'{label}' was longer than {cap} characters and was shortened. Everything else on "
                "the row was kept."
            ),
            value[:200],
        )
    )
    return clip(value, cap)


def _details_plan_year(rows: list[tuple[int, list[Any]]], sheet: str, problems: list[ParseProblem]):
    """The plan year off a Details sheet, READ OUT OF THE RAW VALUES.

    ⚠ THIS IS NOT `xlsx_table.read_details`, AND THAT IS THE POINT. That helper declares — and
    means — `dict[str, str]`: it is written for TEXT rows, and its value picker uses a bare
    truthiness test. Handed value rows it would answer an `int` for a year typed as a number and a
    `datetime` for the cell Excel auto-formatted, both of which its annotation says cannot happen,
    and `read_plan_year` would then be reading a year out of the string "2026-04-01 00:00:00".

    Widening the shared helper was the alternative and was rejected: it is used by the questionnaire
    and the artisan readers, both of which want text and one of which has shipped, and a shared
    function whose return type depends on which row shape it was handed is a contract nobody can
    read off the signature. So the ONE label this sheet carries is matched here, over the values, in
    nine lines that borrow `read_details`' matching rule and nothing else.
    """
    aliases = _DETAIL_KEYS["planYear"]
    for row_number, values in rows[:60]:
        if not values:
            continue
        if norm_header(values[0]) not in aliases:
            continue
        raw = next(
            (v for v in values[1:] if not isinstance(v, FormulaCell) and cell_text(v)),
            None,
        )
        if raw is None:
            # A FORMULA IN THE VALUE CELL IS STILL WORTH SAYING. `read_plan_year` names it, and
            # skipping it here would answer "this workbook does not say which year" for a sheet
            # that visibly does.
            raw = next((v for v in values[1:] if isinstance(v, FormulaCell)), None)
            if raw is None:
                return None
        return read_plan_year(raw, sheet=sheet, row=row_number, problems=problems)
    return None


def _details_filter(rows: list[tuple[int, list[Any]]]) -> str | None:
    """The filter this workbook declares it was exported under, or ``None``.

    THE SHAPE IS `_details_plan_year`'S AND THE REASON IS THE SAME ONE: this reader walks VALUE
    rows, and `xlsx_table.read_details` declares — and means — `dict[str, str]`. Handed value rows
    it would answer whatever openpyxl put in the cell.

    **NO PROBLEM IS EVER FILED FROM HERE.** An unreadable, formula-bearing or absent marker answers
    `None`, which means "this sheet does not say it is partial" — and the upload then proceeds
    exactly as it did before this marker existed. That is the safe direction for a NEW field:
    an office that types over the Details sheet, an older file exported before this row was
    written, or a directory built from scratch in Excel must not start being refused.

    ⚠ **`None` IS NOT A GUARANTEE THAT THE SHEET IS THE WHOLE YEAR**, and the refusal that reads
    it must be worded as what it is — a refusal of a sheet that SAYS it is partial, not a promise
    about every sheet that does not.
    """
    aliases = _DETAIL_KEYS["filter"]
    for _row_number, values in rows[:60]:
        if not values:
            continue
        if norm_header(values[0]) not in aliases:
            continue
        for value in values[1:]:
            if isinstance(value, FormulaCell):
                continue
            text = cell_text(value).strip()
            if text:
                return text[:MAX_FILTER_NOTE_CHARS]
        return None
    return None


def _read_rows(
    rows: list[tuple[int, list[Any]]],
    sheet: str,
    header_row: int,
    columns: dict[int, tuple[str, str | None]],
    result: ParsedAnnualPlan,
) -> None:
    """Walk the data rows under the header. **NEVER DROPS A ROW IN SILENCE.**

    A row is either a :class:`ParsedPlanRow`, or a wholly blank spacer, or a ``ParseProblem`` with
    ``severity="error"``. There is no fourth outcome and no ``continue`` that leaves nothing behind.
    """
    labels = dict(_PRO_FORMA_COLUMNS)
    seen: dict[str, int] = {}

    for row_number, values in rows:
        if row_number <= header_row:
            continue

        def get(role: str, _values: list[Any] = values) -> Any:
            for index, (found, _label) in columns.items():
                if found == role and index < len(_values):
                    return _values[index]
            return None

        if not any(cell_text(v) for v in values):
            # A SPACER ROW IS SKIPPED IN SILENCE AND THAT IS THE ONE SILENT SKIP IN THIS FUNCTION.
            # Ministry directories are full of them — a blank line under each state's block — and a
            # problem per blank line would bury the two problems that matter under forty that do not.
            continue

        if len(result.rows) >= MAX_PLAN_ROWS:
            result.problems.append(
                ParseProblem(
                    sheet,
                    row_number,
                    "error",
                    (
                        f"This sheet has more than {MAX_PLAN_ROWS} rows. The first {MAX_PLAN_ROWS} "
                        f"were imported and everything from row {row_number} down was not. A "
                        "year's directory is 200-300 workshops; if this file really is bigger, "
                        "split it and upload each part."
                    ),
                    None,
                )
            )
            break

        raw_no = cell_text(get("workshopNo"))
        if not raw_no:
            # The other columns' data is named in `value`, so an administrator can find the row in a
            # sheet whose gutter numbers they have re-sorted since.
            trace = ", ".join(t for v in values if (t := cell_text(v)))[:200]
            result.problems.append(
                ParseProblem(
                    sheet,
                    row_number,
                    "error",
                    (
                        "This row has no Workshop No., so there is nothing to file it under and "
                        "nothing to match it to when you upload a correction. Type the ministry's "
                        "reference for it, or delete the row."
                    ),
                    trace or None,
                )
            )
            continue

        workshop_no = raw_no
        if len(workshop_no) > _TEXT_CAPS["workshopNo"]:
            workshop_no = clip(workshop_no, _TEXT_CAPS["workshopNo"])
            result.problems.append(
                ParseProblem(
                    sheet,
                    row_number,
                    "warning",
                    (
                        f"The Workshop No. was longer than {_TEXT_CAPS['workshopNo']} characters "
                        "and was shortened. A workshop code is copied onto the workshop itself, "
                        "which is bounded at that length."
                    ),
                    raw_no[:200],
                )
            )

        key = fold_workshop_no(workshop_no)
        if key in seen:
            # PRE-EMPTING THE ONLY CONSTRAINT AN UPLOAD CAN VIOLATE. `@@unique([planYear,
            # workshopNoKey])` is the one index a sheet can break, and breaking it inside the write
            # transaction would roll the whole directory back over a mistake the uploader could have
            # been told about here. Both Excel rows are named because the uploader has to look at
            # both to decide which is right.
            result.problems.append(
                ParseProblem(
                    sheet,
                    row_number,
                    "error",
                    (
                        f"Workshop No. '{workshop_no}' is also on row {seen[key]} of this sheet. "
                        "The same workshop cannot be planned twice in one year, so this row was "
                        "not imported. Give it its own number or delete it."
                    ),
                    workshop_no,
                )
            )
            continue
        seen[key] = row_number

        entry = ParsedPlanRow(sheetRow=row_number, workshopNo=workshop_no, workshopNoKey=key)

        kind_text = cell_text(get("workshopKind"))
        if kind_text:
            entry.workshopKind = _read_workshop_kind(
                kind_text, sheet=sheet, row=row_number, problems=result.problems
            )

        state_text = cell_text(get("state"))
        if state_text:
            # CLIPPED LIKE EVERY OTHER PROMOTED COLUMN, and it has to be clipped HERE rather than in
            # the loop below because `normalize_state` is what writes it. An unrecognised spelling
            # comes back TRIMMED BUT NOT BOUNDED — that is the function's whole contract, "kept as
            # typed so it can be quoted back" — so a paragraph pasted into the State column would be
            # promoted onto `DesignWorkshop.state` with no stage entry behind it (the FieldSpec is
            # bounded at 80 and `validate_entry` DROPS a refused value, only logging it), and the
            # designer's first stage-1 save would null the column.
            entry.state = _clip_role(
                address.normalize_state(state_text), "state", sheet, row_number, result.problems
            )
            reason = address.state_error(entry.state)
            if reason:
                result.problems.append(
                    ParseProblem(sheet, row_number, "warning", reason, state_text)
                )

        district_text = cell_text(get("district"))
        if district_text:
            entry.district = _clip_role(
                address.normalize_district(entry.state, district_text),
                "district",
                sheet,
                row_number,
                result.problems,
            )
            reason = address.district_error(entry.state, entry.district)
            # A district can only be checked against a state the list knows. When the state itself
            # was the problem it has already been reported, and saying it twice on one row makes the
            # panel read as two mistakes where there is one.
            if reason and not address.state_error(entry.state):
                result.problems.append(
                    ParseProblem(sheet, row_number, "warning", reason, district_text)
                )

        entry.plannedStartDate = read_plan_date(
            get("plannedStartDate"),
            sheet=sheet,
            row=row_number,
            column_label="Start Date",
            problems=result.problems,
        )
        entry.plannedEndDate = read_plan_date(
            get("plannedEndDate"),
            sheet=sheet,
            row=row_number,
            column_label="End Date",
            problems=result.problems,
        )
        if (
            entry.plannedStartDate
            and entry.plannedEndDate
            and entry.plannedEndDate < entry.plannedStartDate
        ):
            # BOTH ARE KEPT. They are what the ministry's document says, and a parser that discards
            # the one it thinks is wrong has destroyed the evidence that the sheet needs correcting.
            result.problems.append(
                ParseProblem(
                    sheet,
                    row_number,
                    "warning",
                    (
                        f"The End Date ({entry.plannedEndDate.isoformat()}) is before the Start "
                        f"Date ({entry.plannedStartDate.isoformat()}). Both were kept as typed — "
                        "correct the sheet and upload it again."
                    ),
                    None,
                )
            )

        for role in ("plannedTitle", "craftName", "clusterName", "venue", "implementingAgency", "sponsor", "notes"):
            text = cell_text(get(role))
            if not text:
                continue
            cap = _TEXT_CAPS[role]
            if len(text) > cap:
                result.problems.append(
                    ParseProblem(
                        sheet,
                        row_number,
                        "warning",
                        (
                            f"'{labels[role]}' was longer than {cap} characters and was shortened. "
                            "Everything else on the row was kept."
                        ),
                        text[:200],
                    )
                )
                text = clip(text, cap)
            setattr(entry, role, text)

        result.rows.append(entry)


# --- Writing ---------------------------------------------------------------------------------------


def _header(ws: Worksheet, row: int, labels: tuple[tuple[str, str], ...]) -> None:
    """The frozen header band. EVERY CELL GOES THROUGH ``xlsx_report._put`` AND NOTHING ELSE.

    ``_put`` is what strips the control characters Excel refuses to open a file containing, what
    holds a value to the 32767-character ceiling, and what makes a leading ``=`` a literal rather
    than a formula. `tests/test_xlsx_report.py` asserts all three over every workbook it knows how
    to build, which is why `test_xlsx_report.py` was taught to build this one.

    (The one call below that is NOT a write — `ws.cell(...).column_letter` — reads a coordinate to
    set a column width. It writes no value, so it is outside that rule rather than an exception to
    it; `artisan_xlsx._header` does the same thing for the same reason.)
    """
    fill = PatternFill(fill_type="solid", start_color=f"FF{_BRAND}")
    required_fill = PatternFill(fill_type="solid", start_color="FF9A3412")
    for index, (key, label) in enumerate(labels, start=1):
        cell = _put(ws, row, index, label)
        cell.font = Font(bold=True, color=_WHITE)
        cell.fill = required_fill if key in _REQUIRED else fill
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.column_dimensions[ws.cell(row=row, column=index).column_letter].width = _WIDTHS.get(
            key, 20
        )
    ws.row_dimensions[row].height = 32
    ws.freeze_panes = f"A{row + 1}"


_DETAILS_NOTE = (
    "Type the financial year this directory is for. Leave every heading on the 'Annual plan' sheet "
    "exactly as you found it — the app finds the columns by their headings, not by where they are. "
    "Uploading a corrected copy of this file updates the workshops it names and creates the ones it "
    "does not; it never makes a second copy of a workshop already in the plan."
)

#: The same note, for a sheet that is only PART of the year, with the sentence that matters most on
#: it. An administrator holding a filtered export has to be told two things the unfiltered note does
#: not say: that the rows they cannot see are still in the plan, and that this file cannot be used
#: to withdraw anything. The refusal in `upload_annual_plan` says the same thing again at the moment
#: it matters; this says it while they are still in Excel, which is where it is cheap to act on.
_DETAILS_NOTE_FILTERED = (
    "This file is a FILTERED VIEW of the year, not the whole directory — the 'Filtered view' row "
    "above says which filter it was taken under. Correcting these rows and uploading this file "
    "updates the workshops it names and creates the ones it does not; the workshops it does not "
    "name are untouched and stay in the plan. It CANNOT be used to withdraw workshops: an upload "
    "of this file with 'withdraw the workshops missing from this sheet' ticked is refused, because "
    "the rows missing from it are the ones the filter hid rather than the ones the ministry "
    "dropped. Export the year with the filters cleared if that is what you meant to do. Leave "
    "every heading on the 'Annual plan' sheet exactly as you found it — the app finds the columns "
    "by their headings, not by where they are."
)


def _details_sheet(ws: Worksheet, *, plan_year: str, filter_note: str | None = None) -> None:
    """The Details sheet: which year's plan this is, whether it is all of it, and a note under both.

    TWO LABELLED ROWS AT MOST, and `_DETAIL_KEYS` has a key for each. A third row added here for a
    reader's benefit ("Issued by", "Uploaded on") becomes a value the parser stores the moment
    somebody adds its label to `_DETAIL_KEYS` — and `xlsx_table.read_details`' own warning records
    what that cost the questionnaire.

    **THE "Filtered view" ROW IS WRITTEN ONLY WHEN THERE IS A FILTER**, so an ordinary whole-year
    export produces the file it always produced and nothing downstream sees a new row. When it IS
    written it is the load-bearing half of the guard set out at `_DETAIL_KEYS["filter"]`: it is what
    lets `upload_annual_plan` tell a corrected copy of the WHOLE year from a corrected copy of forty
    rows of it, which is the difference between withdrawing nothing and withdrawing two hundred and
    sixty workshops. The sentence is also there to be READ — an administrator who opens the file
    sees, beside the year, that it is not the whole year.
    """
    ws.sheet_properties.tabColor = f"FF{_BRAND}"
    label = _put(ws, 1, 1, "Plan year")
    label.font = Font(bold=True)
    _put(ws, 1, 2, plan_year)
    if filter_note:
        marker = _put(ws, 2, 1, "Filtered view")
        marker.font = Font(bold=True)
        value = _put(ws, 2, 2, filter_note[:MAX_FILTER_NOTE_CHARS])
        value.alignment = Alignment(wrap_text=True, vertical="top")
    note = _put(ws, 3, 1, _DETAILS_NOTE if not filter_note else _DETAILS_NOTE_FILTERED)
    note.font = Font(italic=True, color=_GREY)
    note.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=6)
    ws.row_dimensions[3].height = 64
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 28


# ⚠ NOT ONE OF THESE LINES MAY NORMALISE TO A COLUMN ALIAS, AND THE BOLD SUB-HEADINGS ARE WHY THEY
# READ "The Workshop No. column" RATHER THAN "Workshop No.".
#
# `xlsx_table.scan_header` accepts ANY row that names the required role, and a bold sub-heading
# sitting alone in column A is such a row. With the bare spelling, this instructions sheet was a
# second candidate for the data sheet on every upload, and `pick_sheet` said so in a warning — on a
# file this app itself wrote. A warning that fires on a CORRECT file is how an administrator learns
# to ignore the panel that also tells them a real column was skipped. The example block below is one
# piped string per line for the same reason, stated again there.
_HELP_LINES: tuple[tuple[str, str], ...] = (
    ("h", "How to fill this in"),
    (
        "p",
        "One row per planned workshop, on the 'Annual plan' sheet. Workshop No. is the only column "
        "you must fill in; everything else can be added later by uploading a corrected copy of "
        "this same file.",
    ),
    ("b", "The Workshop No. column"),
    (
        "p",
        "The ministry's own reference for the workshop. It is how the app recognises a workshop "
        "when you upload a corrected sheet, so do not renumber workshops between versions. A "
        "serial number (1, 2, 3) is NOT a workshop number and cannot be used: insert one row at "
        "the top and every number below it changes.",
    ),
    ("b", "The two date columns"),
    (
        "p",
        "Format the date columns as a date in Excel and the app never guesses. Typed as plain "
        "text, 03/04/2026 is read as 3 April 2026 — day first — and you are told so on the upload "
        "report, because 3 April and 4 March are both valid readings of that.",
    ),
    ("b", "The State and District columns"),
    (
        "p",
        "Spelled as the Government of India's list spells them. A spelling the app does not "
        "recognise is KEPT exactly as you typed it and reported on the upload report; it is never "
        "silently blanked.",
    ),
    ("b", "The Type of Workshop column"),
    (
        "p",
        "One of: " + ", ".join(ENUMS.get("WORKSHOP_KIND", {}).values()) + ". Anything else is left "
        "blank and reported, and the rest of the row is kept.",
    ),
    ("b", "Correcting the plan"),
    (
        "p",
        "Upload the corrected sheet again. Rows already in the plan are updated, new rows are "
        "added, and a row you have deleted from the sheet is LEFT ALONE unless you tick 'withdraw "
        "the workshops missing from this sheet' on the upload panel. Uploading the same sheet "
        "twice changes nothing at all.",
    ),
    ("b", "A workshop that has already started"),
    (
        "p",
        "A plan row that has been opened as a real design & prototype workshop is still corrected "
        "in the directory — AND THE WORKSHOP ITSELF IS NOT TOUCHED. The upload report names every "
        "such row. Correct the workshop on its own screen.",
    ),
    ("b", "An example"),
)

# ⚠ ONE CELL PER LINE, NOT ONE CELL PER COLUMN, AND THAT IS LOAD-BEARING RATHER THAN A LAYOUT
# CHOICE. Written as separate cells, the first line of this block IS a header row by every rule
# `xlsx_table.scan_header` applies — it names a "Workshop No." column — so this sheet became a
# SECOND candidate for the data sheet, and `pick_sheet` then emitted "More than one sheet looked
# like an annual plan; Ignored: How to fill this in." on every single upload of a file this app
# itself wrote. A warning that fires on a correct file is how an administrator learns to ignore the
# panel that also tells them a column of real data was skipped. Piped into one string, no cell here
# normalises to any alias, this sheet is not a candidate at all, and the example still reads as the
# table it is describing.
_HELP_EXAMPLE: tuple[str, ...] = (
    "Workshop No.  |  Workshop Title  |  Type of Workshop  |  Craft  |  State  |  District  |  "
    "Venue  |  Start Date",
    "DPW/2026/017  |  Kachchh weaving, Bhuj  |  Design & Prototype Development  |  Kachchh "
    "weaving  |  Gujarat  |  Kachchh  |  DIC Hall, Bhuj  |  12/03/2026",
    "DPW/2026/018  |  (left blank)  |  Skill Upgradation  |  Blue pottery  |  Rajasthan  |  "
    "Jaipur  |  Kala Kendra, Jaipur  |  02/04/2026",
)


def _help_sheet(ws: Worksheet) -> None:
    """The instructions, and the worked example that is deliberately NOT on the data sheet.

    A seeded example row on the 'Annual plan' sheet is a row somebody uploads — the argument
    `questionnaire_xlsx.build_pro_forma` makes about its own pro-forma. On this sheet it cannot be
    imported, and the reason is stated on `_HELP_EXAMPLE` rather than assumed here: every line of
    the example is ONE cell of piped text, so no cell on this sheet normalises to a column alias and
    `scan_header` finds no header row on it at all. `tests/test_annual_plan_xlsx.py::
    test_the_instructions_sheet_is_not_even_a_candidate_for_the_data_sheet` pins that.
    """
    ws.sheet_properties.tabColor = _GREY
    ws.column_dimensions["A"].width = 110
    row = 1
    for kind, text in _HELP_LINES:
        cell = _put(ws, row, 1, text)
        if kind == "h":
            cell.font = Font(bold=True, size=14, color=f"FF{_BRAND}")
        elif kind == "b":
            cell.font = Font(bold=True)
        else:
            cell.font = Font(color=_GREY)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[row].height = 42
        row += 1
    row += 1
    for index, line in enumerate(_HELP_EXAMPLE):
        cell = _put(ws, row, 1, line)
        cell.font = Font(bold=index == 0, color=_GREY, name="Consolas", size=9)
        row += 1


def build_annual_plan_pro_forma() -> bytes:
    """The BLANK pro-forma: headings, a Details sheet naming the plan year, and instructions.

    EMPTY UNDER THE HEADINGS, for the reason `questionnaire_xlsx.build_pro_forma` gives about its
    own: a seeded example row is a row somebody uploads, and then the ministry's directory contains
    a workshop called "Example". The worked example lives on the instructions sheet, where it cannot
    be imported.
    """
    wb = Workbook()
    plan = wb.active
    plan.title = SHEET_PLAN
    plan.sheet_properties.tabColor = f"FF{_BRAND}"
    _header(plan, 1, _PRO_FORMA_COLUMNS)
    _details_sheet(wb.create_sheet(title=SHEET_DETAILS), plan_year="")
    _help_sheet(wb.create_sheet(title=SHEET_HELP))
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def build_annual_plan_workbook(
    *, plan_year: int | None, rows: list[dict[str, Any]], filter_note: str | None = None
) -> bytes:
    """The current directory as the same workbook — the download half of edit-in-Excel.

    THE ROUND TRIP IS THE FEATURE: download, correct in Excel, upload the same file. So this emits
    exactly the pro-forma's columns and headings, in the pro-forma's order, and one row per entry in
    the list's own order — which is the server's sort, because the sort is the server's.

    THE TWO DATE COLUMNS ARE WRITTEN AS REAL DATE CELLS AND NOT AS TEXT, and that is the whole
    reason `test_a_generated_workbooks_dates_round_trip_without_an_ambiguity_warning` exists. A file
    this app wrote and this app reads back must produce ZERO ambiguity warnings; text dates would
    produce one per row on every export/import cycle and teach an administrator that the warning
    means nothing.
    """
    wb = Workbook()
    plan = wb.active
    plan.title = SHEET_PLAN
    plan.sheet_properties.tabColor = f"FF{_BRAND}"
    _header(plan, 1, _PRO_FORMA_COLUMNS)

    kinds: dict[str, str] = ENUMS.get("WORKSHOP_KIND", {})
    for offset, row in enumerate(rows, start=2):
        for index, (key, _label) in enumerate(_PRO_FORMA_COLUMNS, start=1):
            value = row.get(key)
            if key == "workshopKind" and value:
                value = kinds.get(str(value), value)
            elif key in ("plannedStartDate", "plannedEndDate") and value is not None:
                value = value.date() if isinstance(value, datetime) else value
            cell = _put(plan, offset, index, value)
            if key in ("plannedStartDate", "plannedEndDate") and value is not None:
                cell.number_format = "DD/MM/YYYY"

    _details_sheet(
        wb.create_sheet(title=SHEET_DETAILS),
        plan_year=plan_year_label(plan_year),
        filter_note=filter_note,
    )
    _help_sheet(wb.create_sheet(title=SHEET_HELP))
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
